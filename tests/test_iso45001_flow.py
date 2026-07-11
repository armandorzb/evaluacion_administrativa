from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage

from municipal_diagnostico import create_app
from municipal_diagnostico.extensions import db
from municipal_diagnostico.iso45001_seed_data import ISO45001_V2_CATALOG_SLUG
from municipal_diagnostico.models import (
    Dependencia,
    Iso45001Apartado,
    Iso45001Asignacion,
    Iso45001Ciclo,
    Iso45001Clausula,
    Iso45001CuestionarioVersion,
    Iso45001DocumentoRequerido,
    Iso45001Evidencia,
    Iso45001EvidenciaDocumental,
    Iso45001Evaluacion,
    Iso45001ObservacionRevision,
    Iso45001Reactivo,
    Iso45001Respuesta,
    Usuario,
)
from municipal_diagnostico.services.iso45001 import (
    ISO45001_OPTION_POINTS,
    ensure_document_controls_for_evaluation,
    save_document_control_payload,
    summarize_iso45001_evaluation,
    upload_document_evidence,
    validate_iso45001_submission,
)
from municipal_diagnostico.services.iso45001_exports import (
    build_iso45001_excel,
    build_iso45001_pdf,
)


class TestConfig:
    SECRET_KEY = "test"
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    UPLOAD_FOLDER = "tests/uploads-iso45001"
    ALLOWED_EXTENSIONS = {"pdf"}
    AUTO_INIT_DATABASE = True
    BOOTSTRAP_ADMIN_EMAIL = None
    BOOTSTRAP_ADMIN_PASSWORD = None
    BOOTSTRAP_ADMIN_NAME = None


def login(client, email: str, password: str = "secret123"):
    return client.post(
        "/auth/login",
        data={"correo": email, "password": password},
        follow_redirects=True,
    )


def build_app(upload_folder: Path | None = None):
    config = TestConfig
    if upload_folder is not None:
        config = type("Iso45001UploadTestConfig", (TestConfig,), {"UPLOAD_FOLDER": str(upload_folder)})

    app = create_app(config)
    with app.app_context():
        dependency_one = Dependencia(nombre="Contraloria", tipo="Administrativa")
        dependency_two = Dependencia(nombre="Secretaria Tecnica", tipo="Administrativa")
        users = [
            Usuario(
                nombre="Admin ISO 45001",
                correo="admin.iso45001@test.local",
                rol="administrador",
                activo=True,
                acceso_diagnostico=False,
                acceso_bienestar=False,
                acceso_iso9001=False,
                acceso_iso45001=True,
            ),
            Usuario(
                nombre="Capturista ISO 45001",
                correo="captura.iso45001@test.local",
                rol="evaluador",
                activo=True,
                acceso_diagnostico=False,
                acceso_bienestar=False,
                acceso_iso9001=False,
                acceso_iso45001=True,
            ),
            Usuario(
                nombre="Revisor ISO 45001",
                correo="revisor.iso45001@test.local",
                rol="revisor",
                activo=True,
                acceso_diagnostico=False,
                acceso_bienestar=False,
                acceso_iso9001=False,
                acceso_iso45001=True,
            ),
            Usuario(
                nombre="Consulta ISO 45001",
                correo="consulta.iso45001@test.local",
                rol="consulta",
                activo=True,
                acceso_diagnostico=False,
                acceso_bienestar=False,
                acceso_iso9001=False,
                acceso_iso45001=True,
            ),
            Usuario(
                nombre="Solo ISO 9001",
                correo="solo.iso9001@test.local",
                rol="consulta",
                activo=True,
                acceso_diagnostico=False,
                acceso_bienestar=False,
                acceso_iso9001=True,
                acceso_iso45001=False,
            ),
            Usuario(
                nombre="Sin modulos ISO",
                correo="sin.iso@test.local",
                rol="consulta",
                activo=True,
                acceso_diagnostico=True,
                acceso_bienestar=False,
                acceso_iso9001=False,
                acceso_iso45001=False,
            ),
        ]
        for user in users:
            user.set_password("secret123")
        db.session.add_all([dependency_one, dependency_two, *users])
        db.session.commit()
        return app, {
            "dependency_one_id": dependency_one.id,
            "dependency_two_id": dependency_two.id,
            "admin_id": users[0].id,
            "evaluator_id": users[1].id,
            "reviewer_id": users[2].id,
            "consultant_id": users[3].id,
            "iso9001_only_id": users[4].id,
        }


def create_iso45001_evaluation(
    dependency_id: int,
    evaluator_id: int,
    reviewer_id: int,
    admin_id: int,
    *,
    cycle_name: str | None = None,
    version_slug: str = ISO45001_V2_CATALOG_SLUG,
) -> int:
    version = Iso45001CuestionarioVersion.query.filter_by(slug=version_slug).one()
    cycle = Iso45001Ciclo(
        nombre=cycle_name or f"Ciclo ISO 45001 {dependency_id}",
        descripcion="Ciclo de prueba",
        estado="activo",
        fecha_inicio=date(2026, 1, 1),
        fecha_cierre=date(2026, 12, 31),
        version=version,
        creado_por_id=admin_id,
    )
    evaluation = Iso45001Evaluacion(
        ciclo=cycle,
        dependencia_id=dependency_id,
        revisor_id=reviewer_id,
        estado="borrador",
    )
    db.session.add_all([cycle, evaluation])
    db.session.flush()
    db.session.add(
        Iso45001Asignacion(
            evaluacion=evaluation,
            usuario_id=evaluator_id,
            tipo="captura",
        )
    )
    ensure_document_controls_for_evaluation(
        evaluation,
        user=db.session.get(Usuario, evaluator_id),
    )
    db.session.commit()
    return evaluation.id


def reactives_for_version(version: Iso45001CuestionarioVersion) -> list[Iso45001Reactivo]:
    return [
        reactive
        for clause in version.clausulas
        for section in clause.apartados
        for reactive in section.reactivos
    ]


def version_reactives(evaluation: Iso45001Evaluacion) -> list[Iso45001Reactivo]:
    return reactives_for_version(evaluation.ciclo.version)


def fill_all_responses(evaluation_id: int, user_id: int, value: str = "no") -> None:
    evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
    responses = {response.reactivo_id: response for response in evaluation.respuestas}
    for reactive in version_reactives(evaluation):
        response = responses.get(reactive.id)
        if response is None:
            response = Iso45001Respuesta(
                evaluacion=evaluation,
                reactivo=reactive,
                usuario_id=user_id,
            )
            db.session.add(response)
        response.calificacion = value
        response.valor = ISO45001_OPTION_POINTS[value]
    summarize_iso45001_evaluation(evaluation)
    db.session.commit()


def mapped_reactive(version: Iso45001CuestionarioVersion) -> Iso45001Reactivo:
    reactive = next(
        (
            reactive
            for clause in version.clausulas
            for section in clause.apartados
            for reactive in section.reactivos
            if reactive.documentos_requeridos
        ),
        None,
    )
    assert reactive is not None
    assert reactive.documentos_requeridos
    return reactive


def mark_all_document_controls_no(evaluation_id: int, user_id: int) -> None:
    """Evaluate every v2 documentary control as a documented gap without files."""

    evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
    user = db.session.get(Usuario, user_id)
    controls = ensure_document_controls_for_evaluation(evaluation, user=user)
    assert len(controls) == 37
    for assessment in controls:
        result = save_document_control_payload(
            evaluation,
            assessment.id,
            {
                "punto_ids": [],
                "evidencia_ids": [],
                "observacion": f"Brecha documentada para {assessment.control.codigo}.",
            },
            user=user,
        )
        assert result["ok"] is True
    summarize_iso45001_evaluation(evaluation)
    db.session.commit()


def all_text(worksheet) -> list[str]:
    return [
        value
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if isinstance(value, str)
    ]


def test_iso45001_catalog_has_fixed_matrix_documents_and_amd1_controls():
    app, _ids = build_app()

    with app.app_context():
        v1 = Iso45001CuestionarioVersion.query.filter_by(
            slug="iso45001_2018_amd1_2024_diagnostico_v1"
        ).one()
        version = Iso45001CuestionarioVersion.query.filter_by(
            slug=ISO45001_V2_CATALOG_SLUG
        ).one()
        assert {v1.slug, version.slug} == {
            "iso45001_2018_amd1_2024_diagnostico_v1",
            ISO45001_V2_CATALOG_SLUG,
        }
        assert len(v1.clausulas) == 7
        assert len(v1.documentos_obligatorios) == 31
        assert len(v1.controles_evidencia) == 0
        assert len(version.clausulas) == 7
        assert sum(len(clause.apartados) for clause in version.clausulas) == 40
        assert len(reactives_for_version(version)) == 308
        assert len(version.documentos_obligatorios) == 31
        assert len(version.controles_evidencia) == 37
        assert sum(
            control.tipo == "documento_normativo"
            for control in version.controles_evidencia
        ) == 31
        assert sum(
            control.tipo == "grupo_complementario"
            for control in version.controles_evidencia
        ) == 6
        assert Iso45001Clausula.query.count() == 14
        assert Iso45001Apartado.query.count() == 80
        assert Iso45001Reactivo.query.count() == 616
        assert Iso45001DocumentoRequerido.query.count() == 62

        counts = {
            int(clause.numero): sum(len(section.reactivos) for section in clause.apartados)
            for clause in version.clausulas
        }
        assert counts == {4: 27, 5: 39, 6: 64, 7: 57, 8: 55, 9: 42, 10: 24}

        amendments = {
            (reactive.codigo, reactive.apartado.codigo)
            for reactive in reactives_for_version(version)
            if reactive.es_enmienda_2024
        }
        assert amendments == {("R-307", "4.1"), ("R-308", "4.2")}

        documents = sorted(version.documentos_obligatorios, key=lambda document: document.orden)
        assert documents[0].codigo == "D-01"
        assert documents[-1].codigo == "D-31"
        assert {"Mantener", "Conservar", "Mantener y conservar"}.issubset(
            {document.clasificacion for document in documents}
        )
        assert any(reactive.documentos_requeridos for reactive in reactives_for_version(version))

        controls = {control.codigo: control for control in version.controles_evidencia}
        assert set(controls) == {
            *(f"D-{number:02d}" for number in range(1, 32)),
            *(f"G-{number:02d}" for number in range(1, 7)),
        }
        assert all(control.puntos for control in controls.values())

        def reactive_codes(control_code: str) -> list[str]:
            return [reactive.codigo for reactive in controls[control_code].reactivos]

        assert reactive_codes("D-13") == [f"R-{number:03d}" for number in range(166, 171)]
        assert reactive_codes("D-23") == [
            *(f"R-{number:03d}" for number in range(241, 248)),
            "R-249",
            "R-250",
        ]
        assert reactive_codes("D-24") == ["R-248", "R-250"]
        assert reactive_codes("D-26") == [
            *(f"R-{number:03d}" for number in range(263, 268)),
            "R-270",
        ]
        assert reactive_codes("D-27") == [f"R-{number:03d}" for number in range(268, 271)]
        assert reactive_codes("D-29") == ["R-287", "R-288", "R-290", "R-291", "R-292", "R-298"]
        assert reactive_codes("D-30") == [
            "R-289",
            *(f"R-{number:03d}" for number in range(293, 299)),
        ]
        assert reactive_codes("G-01") == [f"R-{number:03d}" for number in range(7, 14)]
        assert reactive_codes("G-02") == [f"R-{number:03d}" for number in range(53, 65)]
        assert reactive_codes("G-03") == [f"R-{number:03d}" for number in range(71, 85)]
        assert reactive_codes("G-04") == [f"R-{number:03d}" for number in range(156, 161)]
        assert reactive_codes("G-05") == [f"R-{number:03d}" for number in range(161, 166)]
        assert reactive_codes("G-06") == [f"R-{number:03d}" for number in range(258, 263)]


def test_iso45001_rejects_na_and_scores_every_reactive():
    app, ids = build_app()
    client = app.test_client()

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        section = evaluation.ciclo.version.clausulas[0].apartados[0]
        reactive_id = section.reactivos[0].id

    login(client, "captura.iso45001@test.local")
    autosave = client.post(
        f"/iso45001/evaluaciones/{evaluation_id}/apartados/{section.id}/autosave",
        json={
            "responses": [
                {
                    "reactivo_id": reactive_id,
                    "calificacion": "na",
                    "observacion": "No aplica",
                }
            ]
        },
    )
    assert autosave.status_code == 400
    assert autosave.get_json()["ok"] is False

    with app.app_context():
        assert Iso45001Respuesta.query.filter_by(evaluacion_id=evaluation_id).count() == 0
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        user = db.session.get(Usuario, ids["evaluator_id"])
        reactives = version_reactives(evaluation)[:3]
        for reactive, value in zip(reactives, ["si", "parcial", "no"]):
            db.session.add(
                Iso45001Respuesta(
                    evaluacion=evaluation,
                    reactivo=reactive,
                    usuario=user,
                    calificacion=value,
                    valor=ISO45001_OPTION_POINTS[value],
                )
            )
        db.session.flush()

        summary = summarize_iso45001_evaluation(evaluation)
        assert ISO45001_OPTION_POINTS == {"no": 0, "parcial": 1, "si": 2}
        assert summary["total_questions"] == 308
        assert summary["applicable_questions"] == 308
        assert summary["na_questions"] == 0
        assert summary["answered_questions"] == 3
        assert summary["points"] == 3
        assert summary["percent"] == round((3 / (308 * 2)) * 100, 2)
        assert summary["completion"] == round((3 / 308) * 100, 2)
        assert summary["response_counts"] == {
            "no": 1,
            "parcial": 1,
            "si": 1,
            "sin_respuesta": 305,
        }
        assert summary["uses_document_control_coverage"] is True
        assert summary["document_evidence_stats"]["total_controls"] == 37


def test_iso45001_v1_evidence_gate_is_preserved_for_historical_cycles():
    app, ids = build_app()

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
            version_slug="iso45001_2018_amd1_2024_diagnostico_v1",
        )
        fill_all_responses(evaluation_id, ids["evaluator_id"], "no")

        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        assert evaluation.ciclo.version.slug == "iso45001_2018_amd1_2024_diagnostico_v1"
        assert summarize_iso45001_evaluation(evaluation)["document_controls"] == []
        reactive = mapped_reactive(evaluation.ciclo.version)
        response = Iso45001Respuesta.query.filter_by(
            evaluacion_id=evaluation_id,
            reactivo_id=reactive.id,
        ).one()

        no_validation = validate_iso45001_submission(evaluation)
        assert no_validation["ok"] is True
        assert no_validation["missing_evidence"] == []

        response.calificacion = "parcial"
        response.valor = ISO45001_OPTION_POINTS["parcial"]
        db.session.flush()
        partial_validation = validate_iso45001_submission(evaluation)
        assert partial_validation["ok"] is False
        assert [row["reactivo"].id for row in partial_validation["missing_evidence"]] == [reactive.id]

        response.calificacion = "si"
        response.valor = ISO45001_OPTION_POINTS["si"]
        db.session.flush()
        si_validation = validate_iso45001_submission(evaluation)
        assert si_validation["ok"] is False
        assert [row["reactivo"].id for row in si_validation["missing_evidence"]] == [reactive.id]

        db.session.add(
            Iso45001Evidencia(
                respuesta=response,
                usuario_id=ids["evaluator_id"],
                archivo_nombre_original="evidencia.pdf",
                archivo_guardado="iso45001/evidencia.pdf",
                mime_type="application/pdf",
                tamano_bytes=12,
            )
        )
        db.session.flush()
        assert validate_iso45001_submission(evaluation)["ok"] is True


def test_iso45001_v2_document_controls_allow_gap_notes_and_require_reusable_files(tmp_path):
    app, ids = build_app(tmp_path)

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )
        fill_all_responses(evaluation_id, ids["evaluator_id"], "no")
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        assert evaluation.ciclo.version.slug == ISO45001_V2_CATALOG_SLUG
        assert len(evaluation.controles_evidencia) == 37

        initial_validation = validate_iso45001_submission(evaluation)
        assert initial_validation["missing_evidence"] == []
        assert len(initial_validation["missing_document_controls"]) == 37

        mark_all_document_controls_no(evaluation_id, ids["evaluator_id"])
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        no_validation = validate_iso45001_submission(evaluation)
        assert no_validation["ok"] is True
        assert no_validation["missing_control_observations"] == []
        assert no_validation["missing_control_evidence"] == []

        # A response can be Parcial without adding a duplicate reactive file:
        # evidence is assessed only through its documented control.
        reactive = mapped_reactive(evaluation.ciclo.version)
        response = Iso45001Respuesta.query.filter_by(
            evaluacion_id=evaluation_id,
            reactivo_id=reactive.id,
        ).one()
        response.calificacion = "parcial"
        response.valor = ISO45001_OPTION_POINTS["parcial"]
        db.session.flush()
        assert validate_iso45001_submission(evaluation)["ok"] is True

        controls = {
            item.control.codigo: item
            for item in ensure_document_controls_for_evaluation(evaluation)
        }
        first = controls["D-01"]
        second = controls["D-02"]
        user = db.session.get(Usuario, ids["evaluator_id"])

        partial_without_file = save_document_control_payload(
            evaluation,
            first.id,
            {
                "punto_ids": [first.control.puntos[0].id],
                "evidencia_ids": [],
                "observacion": "Cobertura parcial pendiente de soportar.",
            },
            user=user,
        )
        assert partial_without_file["ok"] is True
        partial_validation = validate_iso45001_submission(evaluation)
        assert {row["codigo"] for row in partial_validation["missing_control_evidence"]} == {"D-01"}

        all_points_without_file = save_document_control_payload(
            evaluation,
            first.id,
            {
                "punto_ids": [point.id for point in first.control.puntos],
                "evidencia_ids": [],
                "observacion": "Todos los puntos declarados, sin archivo todavía.",
            },
            user=user,
        )
        assert all_points_without_file["ok"] is True
        complete_without_file = validate_iso45001_submission(evaluation)
        assert {row["codigo"] for row in complete_without_file["missing_control_evidence"]} == {"D-01"}

        second_partial = save_document_control_payload(
            evaluation,
            second.id,
            {
                "punto_ids": [second.control.puntos[0].id],
                "evidencia_ids": [],
                "observacion": "El mismo soporte cubrirá este control.",
            },
            user=user,
        )
        assert second_partial["ok"] is True
        uploaded = upload_document_evidence(
            evaluation,
            [
                FileStorage(
                    stream=BytesIO(b"evidencia documental reutilizable"),
                    filename="control-compartido.pdf",
                    content_type="application/pdf",
                )
            ],
            [first.id, second.id],
            user=user,
        )
        assert uploaded["ok"] is True
        assert uploaded["count"] == 1
        db.session.commit()

        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        final_validation = validate_iso45001_submission(evaluation)
        assert final_validation["ok"] is True
        summary = summarize_iso45001_evaluation(evaluation)
        assert summary["document_evidence_stats"]["file_count"] == 1
        assert summary["document_evidence_stats"]["sustained_controls"] == 2
        assert {
            control["codigo"]
            for control in summary["document_controls"]
            if control["files_count"]
        } == {"D-01", "D-02"}

        # The same file cannot be attached to a control from another evaluation.
        other_evaluation_id = create_iso45001_evaluation(
            ids["dependency_two_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )
        other_evaluation = db.session.get(Iso45001Evaluacion, other_evaluation_id)
        other_control = ensure_document_controls_for_evaluation(other_evaluation)[0]
        isolated = save_document_control_payload(
            other_evaluation,
            other_control.id,
            {
                "punto_ids": [other_control.control.puntos[0].id],
                "evidencia_ids": [uploaded["evidences"][0].id],
                "observacion": "No debe aceptar evidencia de otra evaluación.",
            },
            user=user,
        )
        assert isolated["ok"] is False
        assert "esta evaluación" in isolated["error"].lower()


def test_iso45001_capture_review_return_and_close_flow_requires_complete_capture():
    app, ids = build_app()
    client = app.test_client()

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )

    login(client, "captura.iso45001@test.local")
    incomplete_submit = client.post(
        f"/iso45001/evaluaciones/{evaluation_id}/enviar",
        follow_redirects=True,
    )
    assert incomplete_submit.status_code == 200
    assert "Debes responder todos los reactivos" in incomplete_submit.get_data(as_text=True)
    with app.app_context():
        assert db.session.get(Iso45001Evaluacion, evaluation_id).estado == "borrador"
        fill_all_responses(evaluation_id, ids["evaluator_id"], "no")
        mark_all_document_controls_no(evaluation_id, ids["evaluator_id"])

    submit = client.post(f"/iso45001/evaluaciones/{evaluation_id}/enviar", follow_redirects=True)
    assert submit.status_code == 200
    with app.app_context():
        assert db.session.get(Iso45001Evaluacion, evaluation_id).estado == "en_revision"

    client.get("/auth/logout", follow_redirects=True)
    login(client, "revisor.iso45001@test.local")
    returned = client.post(
        f"/iso45001/evaluaciones/{evaluation_id}/revision",
        data={"action": "return", "comentario": "Completar observaciones."},
        follow_redirects=True,
    )
    assert returned.status_code == 200
    with app.app_context():
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        assert evaluation.estado == "devuelta"
        section = evaluation.ciclo.version.clausulas[0].apartados[0]
        reactive_id = section.reactivos[0].id
        section_id = section.id

    client.get("/auth/logout", follow_redirects=True)
    login(client, "captura.iso45001@test.local")
    autosave = client.post(
        f"/iso45001/evaluaciones/{evaluation_id}/apartados/{section_id}/autosave",
        json={
            "responses": [
                {
                    "reactivo_id": reactive_id,
                    "calificacion": "no",
                    "observacion": "Confirmado tras la devolución.",
                }
            ]
        },
    )
    assert autosave.status_code == 200
    assert autosave.get_json()["ok"] is True
    with app.app_context():
        assert db.session.get(Iso45001Evaluacion, evaluation_id).estado == "en_captura"

    submit_again = client.post(f"/iso45001/evaluaciones/{evaluation_id}/enviar", follow_redirects=True)
    assert submit_again.status_code == 200

    client.get("/auth/logout", follow_redirects=True)
    login(client, "revisor.iso45001@test.local")
    closed = client.post(
        f"/iso45001/evaluaciones/{evaluation_id}/revision",
        data={"action": "close", "comentario": "Cierre oficial."},
        follow_redirects=True,
    )
    assert closed.status_code == 200
    with app.app_context():
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        assert evaluation.estado == "cerrada"
        assert evaluation.cerrada_at is not None
        assert [
            observation.accion
            for observation in Iso45001ObservacionRevision.query.filter_by(evaluacion_id=evaluation_id)
            .order_by(Iso45001ObservacionRevision.id)
            .all()
        ] == ["return", "close"]


def test_iso45001_access_is_isolated_from_iso9001():
    app, _ids = build_app()
    client = app.test_client()

    login(client, "solo.iso9001@test.local")
    assert client.get("/iso45001/").status_code == 403

    client.get("/auth/logout", follow_redirects=True)
    login(client, "captura.iso45001@test.local")
    assert client.get("/iso45001/").status_code == 200
    assert client.get("/iso9001/").status_code == 403


def test_iso45001_reports_include_required_sheets_findings_and_disclaimer():
    app, ids = build_app()

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )
        fill_all_responses(evaluation_id, ids["evaluator_id"], "no")
        mark_all_document_controls_no(evaluation_id, ids["evaluator_id"])
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)

        pdf = build_iso45001_pdf(evaluation)
        assert pdf.getvalue().startswith(b"%PDF")

        workbook = load_workbook(BytesIO(build_iso45001_excel(evaluation).getvalue()))
        assert workbook.sheetnames == [
            "Resumen",
            "Matriz",
            "Documentos requeridos",
            "Hallazgos",
            "Guía",
            "Trazabilidad documental",
        ]
        assert "R-001" in all_text(workbook["Matriz"])
        assert "D-01" in all_text(workbook["Documentos requeridos"])
        assert "G-01" in all_text(workbook["Documentos requeridos"])
        assert "Archivo" in all_text(workbook["Trazabilidad documental"])
        assert "Brecha prioritaria" in all_text(workbook["Hallazgos"])
        guide_text = " ".join(all_text(workbook["Guía"])).lower()
        assert "diagnóstico de preparación" in guide_text
        assert "no constituye certificación" in guide_text


def test_iso45001_document_evidence_download_requires_view_permission(tmp_path):
    app, ids = build_app(tmp_path)
    client = app.test_client()

    with app.app_context():
        evaluation_id = create_iso45001_evaluation(
            ids["dependency_one_id"],
            ids["evaluator_id"],
            ids["reviewer_id"],
            ids["admin_id"],
        )
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)

        relative_path = (
            Path("iso45001/evaluaciones")
            / str(evaluation_id)
            / "documentacion"
            / "evidencia.pdf"
        )
        absolute_path = tmp_path / relative_path
        absolute_path.parent.mkdir(parents=True)
        absolute_path.write_bytes(b"evidencia autorizada")
        evidence = Iso45001EvidenciaDocumental(
            evaluacion=evaluation,
            usuario_id=ids["evaluator_id"],
            archivo_nombre_original="evidencia.pdf",
            archivo_guardado=str(relative_path).replace("\\", "/"),
            mime_type="application/pdf",
            tamano_bytes=20,
        )
        db.session.add(evidence)
        db.session.commit()
        evidence_id = evidence.id

    login(client, "consulta.iso45001@test.local")
    assert client.get(f"/iso45001/documentacion/evidencias/{evidence_id}/descargar").status_code == 403

    with app.app_context():
        evaluation = db.session.get(Iso45001Evaluacion, evaluation_id)
        evaluation.estado = "cerrada"
        db.session.commit()

    allowed = client.get(f"/iso45001/documentacion/evidencias/{evidence_id}/descargar")
    assert allowed.status_code == 200
    assert allowed.data == b"evidencia autorizada"
