from __future__ import annotations

import math
import os
from html import escape
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import CondPageBreak, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from municipal_diagnostico.services.iso45001 import (
    ISO45001_OPTION_LABELS,
    format_iso45001_datetime,
    maturity_label,
    summarize_iso45001_evaluation,
)
from municipal_diagnostico.timeutils import utcnow


EXCEL_COLORS = {
    "primary": "1F6F78",
    "soft": "E7F5F4",
    "line": "BCD7D4",
    "white": "FFFFFF",
    "warning": "FFF3CD",
    "danger": "FCE4E4",
    "success": "E6F2E7",
}

PDF_THEME = {
    "primary": colors.HexColor("#1F6F78"),
    "primary_soft": colors.HexColor("#E7F5F4"),
    "primary_dark": colors.HexColor("#145058"),
    "line": colors.HexColor("#BCD7D4"),
    "muted": colors.HexColor("#59636A"),
    "red": colors.HexColor("#C62828"),
    "red_soft": colors.HexColor("#FCE4E4"),
    "yellow": colors.HexColor("#E0A100"),
    "yellow_soft": colors.HexColor("#FFF3CD"),
    "green": colors.HexColor("#2E7D32"),
    "green_soft": colors.HexColor("#E6F2E7"),
    "gray": colors.HexColor("#9AA6AB"),
    "gray_soft": colors.HexColor("#F1F4F5"),
    "ink": colors.HexColor("#193136"),
    "paper": colors.HexColor("#F7FBFB"),
}

MATURITY_GUIDE = [
    ("0%", "Nivel 0 - No iniciado"),
    ("1-20%", "Nivel 1 - Inicial"),
    ("21-40%", "Nivel 2 - En desarrollo"),
    ("41-60%", "Nivel 3 - Definido"),
    ("61-80%", "Nivel 4 - Gestionado"),
    ("81-100%", "Nivel 5 - Optimizado"),
]

CLAUSE_GUIDANCE = {
    "4": "Contexto, partes interesadas, alcance y sistema de gestión de SST.",
    "5": "Liderazgo, política, responsabilidades y participación de trabajadores.",
    "6": "Peligros, riesgos, oportunidades, obligaciones y objetivos de SST.",
    "7": "Recursos, competencia, comunicación e información documentada.",
    "8": "Control operacional, jerarquía de controles, contratistas y emergencias.",
    "9": "Medición, evaluación de cumplimiento, auditoría y revisión directiva.",
    "10": "Incidentes, acciones correctivas y mejora continua.",
}

CLAUSE_REPORT_GUIDANCE = {
    "4": {
        "title": "Contexto de la organizaci\u00f3n",
        "focus": "contexto interno y externo, partes interesadas, alcance y procesos del SGSST",
        "support": "Comprender los factores que influyen en el SGSST, las necesidades de trabajadores y otras partes interesadas, el alcance y los procesos necesarios.",
        "evidence": "An\u00e1lisis de contexto, matriz de partes interesadas, alcance documentado, mapa de procesos y revisi\u00f3n de cuestiones clim\u00e1ticas pertinentes.",
    },
    "5": {
        "title": "Liderazgo y participaci\u00f3n de los trabajadores",
        "focus": "compromiso directivo, pol\u00edtica, roles y consulta/participaci\u00f3n de los trabajadores",
        "support": "Demostrar liderazgo visible, responsabilidades claras y mecanismos eficaces para consultar y permitir la participaci\u00f3n de trabajadores.",
        "evidence": "Pol\u00edtica de SST comunicada, nombramientos, minutas de comit\u00e9, consultas, capacitaci\u00f3n y registros de participaci\u00f3n.",
    },
    "6": {
        "title": "Planificaci\u00f3n",
        "focus": "peligros, riesgos, oportunidades, requisitos legales, objetivos y planes de SST",
        "support": "Planear acciones para eliminar peligros, reducir riesgos, cumplir obligaciones y alcanzar objetivos medibles de SST.",
        "evidence": "Identificaci\u00f3n de peligros, evaluaciones de riesgo, matriz legal, objetivos, indicadores, responsables y planes de trabajo.",
    },
    "7": {
        "title": "Apoyo",
        "focus": "recursos, competencia, conciencia, comunicaci\u00f3n e informaci\u00f3n documentada",
        "support": "Asegurar recursos y competencias suficientes, comunicaciones oportunas y control de la informaci\u00f3n documentada del SGSST.",
        "evidence": "Perfiles, capacitaci\u00f3n, evaluaciones de competencia, comunicaciones, documentos vigentes y controles de registros.",
    },
    "8": {
        "title": "Operaci\u00f3n",
        "focus": "controles operacionales, jerarqu\u00eda de controles, cambios, compras, contratistas y emergencias",
        "support": "Planear y controlar las actividades de SST aplicando la jerarqu\u00eda de controles y gestionando cambios, contratistas y situaciones de emergencia.",
        "evidence": "Procedimientos operativos, permisos, inspecciones, controles de contratistas, gesti\u00f3n del cambio, simulacros y planes de emergencia.",
    },
    "9": {
        "title": "Evaluaci\u00f3n del desempe\u00f1o",
        "focus": "seguimiento, medici\u00f3n, cumplimiento, auditor\u00edas internas y revisi\u00f3n directiva",
        "support": "Medir el desempe\u00f1o del SGSST, evaluar el cumplimiento, auditar y revisar los resultados para orientar decisiones.",
        "evidence": "Indicadores, registros de medici\u00f3n y calibraci\u00f3n, evaluaciones legales, programas de auditor\u00eda, hallazgos y actas directivas.",
    },
    "10": {
        "title": "Mejora",
        "focus": "incidentes, no conformidades, acciones correctivas y mejora continua",
        "support": "Gestionar incidentes y no conformidades, investigar causas y verificar que las acciones correctivas y mejoras sean eficaces.",
        "evidence": "Reportes de incidentes, an\u00e1lisis de causa, acciones correctivas, verificaciones de eficacia, tendencias y lecciones aprendidas.",
    },
}

RESPONSE_ORDER = [
    ("no", "No", PDF_THEME["red"]),
    ("parcial", "Parcial", PDF_THEME["yellow"]),
    ("si", "S\u00ed", PDF_THEME["green"]),
    ("sin_respuesta", "Sin respuesta", PDF_THEME["gray"]),
]

MATURITY_REPORT_GUIDE = [
    ("0%", "Nivel 0 - No iniciado", "red"),
    ("1-20%", "Nivel 1 - Inicial", "red"),
    ("21-40%", "Nivel 2 - En desarrollo", "red"),
    ("41-60%", "Nivel 3 - Definido", "yellow"),
    ("61-80%", "Nivel 4 - Gestionado", "yellow"),
    ("81-100%", "Nivel 5 - Optimizado", "green"),
]

ISO45001_PDF_FOOTER = "Direcci\u00f3n de Recursos Humanos - Ayuntamiento de Hermosillo"


def _evaluation_dependency_name(evaluation) -> str:
    dependency = getattr(evaluation, "dependencia", None)
    return getattr(dependency, "nombre", None) or "Sin dependencia"


def _evaluation_unit_name(evaluation) -> str:
    unit_name = getattr(evaluation, "unidad_administrativa_nombre", None)
    if unit_name:
        return str(unit_name)
    area = getattr(evaluation, "area", None)
    if area is not None and getattr(area, "nombre", None):
        return str(area.nombre)
    return "Alcance hist\u00f3rico por dependencia"


def _evaluation_scope_name(evaluation) -> str:
    return str(getattr(evaluation, "alcance_nombre", None) or _evaluation_unit_name(evaluation))


def _evaluation_scope_description(evaluation) -> str:
    description = getattr(evaluation, "alcance_descripcion", None)
    if description:
        return str(description)
    dependency_name = _evaluation_dependency_name(evaluation)
    if getattr(evaluation, "area", None) is None:
        return f"{dependency_name} | alcance hist\u00f3rico por dependencia"
    return f"{dependency_name} | {_evaluation_unit_name(evaluation)}"


def build_iso45001_excel(evaluation) -> BytesIO:
    summary = summarize_iso45001_evaluation(evaluation)
    uses_document_controls = _uses_document_control_coverage(summary)
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    _write_title(
        summary_sheet,
        "Diagnóstico ISO 45001:2018 + Amd. 1:2024",
        _evaluation_scope_description(evaluation),
        8,
    )
    summary_sheet.append([])
    summary_sheet.append(["Dependencia", _evaluation_dependency_name(evaluation)])
    summary_sheet.append(["Unidad administrativa", _evaluation_unit_name(evaluation)])
    summary_sheet.append(["Estado", summary["state_label"]])
    summary_sheet.append(["Ciclo", evaluation.ciclo.nombre])
    summary_sheet.append(["Avance", _percent_display(summary["completion"])])
    summary_sheet.append(["Cumplimiento", _percent_display(summary["percent"])])
    summary_sheet.append(["Madurez", summary["maturity_label"]])
    if uses_document_controls:
        document_stats = summary["document_evidence_stats"]
        summary_sheet.append(["Archivos documentales únicos", document_stats["file_count"]])
        summary_sheet.append(
            [
                "Controles documentales evaluados",
                f"{document_stats['evaluated_controls']}/{document_stats['total_controls']}",
            ]
        )
        summary_sheet.append(
            [
                "Puntos documentales cubiertos",
                f"{document_stats['covered_points']}/{document_stats['total_points']}",
            ]
        )
        summary_sheet.append(["Cobertura documental", _percent_display(document_stats["percent"])])
    else:
        summary_sheet.append(["Evidencias", summary["evidence_count"]])
        summary_sheet.append(["Cobertura de evidencia requerida", _percent_display(summary["evidence_coverage"]["percent"])])
    summary_sheet.append([])
    summary_sheet.append(
        [
            "Cláusula",
            "Nombre",
            "Reactivos",
            "Respondidos",
            "Puntos",
            "% Cumpl.",
            "Madurez",
            "Cobertura documental" if uses_document_controls else "Evidencias",
        ]
    )
    header_row = summary_sheet.max_row
    for clause in summary["clauses"]:
        clause_document_coverage = clause.get("document_control_coverage")
        clause_support = (
            (
                f"{clause_document_coverage['sustained_controls']}/"
                f"{clause_document_coverage['total_controls']}"
                if clause_document_coverage["total_controls"]
                else "-"
            )
            if uses_document_controls and clause_document_coverage is not None
            else clause.get("evidence_count", 0)
        )
        summary_sheet.append(
            [
                clause["numero"],
                clause["nombre"],
                clause["total"],
                clause["answered"],
                clause["points"],
                _excel_percent(clause["percent"]),
                clause["maturity_label"],
                clause_support,
            ]
        )
    _style_table(summary_sheet, header_row)
    _set_widths(summary_sheet, {"A": 12, "B": 42, "C": 12, "D": 14, "E": 11, "F": 14, "G": 28, "H": 14})

    matrix = workbook.create_sheet("Matriz")
    _write_title(
        matrix,
        "Matriz de diagnóstico ISO 45001",
        f"{_evaluation_scope_description(evaluation)} | {evaluation.ciclo.nombre}",
        12,
    )
    matrix.append([])
    matrix.append(
        [
            "Cláusula",
            "Apartado",
            "ID",
            "Reactivo",
            "Calificación",
            "Puntos",
            "Resultado",
            "Información documentada requerida",
            "Evidencia sugerida",
            "Criterio de conformidad",
            "Observación / hallazgo",
            "Archivos",
        ]
    )
    matrix_header = matrix.max_row
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                reactive = row["reactivo"]
                matrix.append(
                    [
                        clause["numero"],
                        section["codigo"],
                        _reactive_identifier(reactive),
                        reactive.texto,
                        row["selected_label"],
                        row["points"] if row["answered"] else "",
                        _finding_label(row),
                        _control_names(row) if uses_document_controls else _document_names(row),
                        reactive.evidencia_sugerida,
                        reactive.criterio_idoneidad,
                        row["observacion"],
                        ", ".join(evidence.archivo_nombre_original for evidence in row["evidence"]) or "Sin archivos",
                    ]
                )
    _style_table(matrix, matrix_header)
    _set_widths(
        matrix,
        {
            "A": 11,
            "B": 13,
            "C": 12,
            "D": 62,
            "E": 15,
            "F": 10,
            "G": 22,
            "H": 42,
            "I": 42,
            "J": 48,
            "K": 42,
            "L": 28,
        },
    )

    documents = workbook.create_sheet("Documentos requeridos")
    _write_title(
        documents,
        "Información documentada requerida",
        _evaluation_scope_description(evaluation),
        10 if uses_document_controls else 9,
    )
    documents.append([])
    documents.append(
        (
            [
                "Código",
                "Apartado",
                "Tipo",
                "Documento / registro",
                "Contenido mínimo esperado",
                "Reactivos relacionados",
                "Evidencia cargada",
                "Cobertura",
                "Observación",
            ]
            if not uses_document_controls
            else [
                "Código",
                "Apartado",
                "Tipo",
                "Control documental / evidencia",
                "Puntos cubiertos",
                "Estado",
                "Archivos únicos vinculados",
                "Reactivos relacionados",
                "Observación / brecha",
                "Contenido mínimo esperado",
            ]
        )
    )
    documents_header = documents.max_row
    if uses_document_controls:
        for control in summary["document_controls"]:
            documents.append(
                [
                    control["codigo"],
                    control["apartado"],
                    control["clasificacion"],
                    control["nombre"],
                    f"{control['puntos_cubiertos']}/{control['total_puntos']}",
                    control["estado_label"],
                    ", ".join(item["archivo_nombre_original"] for item in control["evidencias"]) or "Sin archivos",
                    ", ".join(item["codigo"] for item in control["reactivos"]) or "-",
                    control["observacion"] or "Sin observación",
                    control["contenido_minimo"],
                ]
            )
    else:
        for document, document_rows in _document_rows(summary):
            evidence_rows = [row for row in document_rows if row["evidence"]]
            selected_rows = [row for row in document_rows if row["selected"] in {"si", "parcial"}]
            documents.append(
                [
                    document.codigo,
                    document.apartado_codigo,
                    document.tipo,
                    document.nombre,
                    document.contenido_minimo,
                    ", ".join(_reactive_identifier(row["reactivo"]) for row in document_rows) or "-",
                    sum(len(row["evidence"]) for row in document_rows),
                    _excel_percent((len(evidence_rows) / len(selected_rows) * 100) if selected_rows else None),
                    _document_observation(document_rows),
                ]
            )
    _style_table(documents, documents_header)
    _set_widths(
        documents,
        (
            {"A": 11, "B": 13, "C": 22, "D": 42, "E": 54, "F": 42, "G": 17, "H": 15, "I": 40}
            if not uses_document_controls
            else {"A": 11, "B": 13, "C": 20, "D": 38, "E": 17, "F": 18, "G": 30, "H": 48, "I": 42, "J": 54}
        ),
    )

    findings = workbook.create_sheet("Hallazgos")
    _write_title(findings, "Hallazgos derivados", _evaluation_scope_description(evaluation), 10)
    findings.append([])
    findings.append(
        [
            "Prioridad",
            "Cláusula",
            "Apartado",
            "ID",
            "Reactivo",
            "Calificación",
            "Evidencia sugerida",
            "Controles relacionados" if uses_document_controls else "Documentos relacionados",
            "Observación",
            "Archivos",
        ]
    )
    findings_header = findings.max_row
    for clause, section, row in _findings(summary):
        reactive = row["reactivo"]
        control_files = (
            _document_control_file_names(
                _document_controls_for_reactive(summary["document_controls"], reactive.id)
            )
            if uses_document_controls
            else ""
        )
        findings.append(
            [
                _finding_label(row),
                clause["numero"],
                section["codigo"],
                _reactive_identifier(reactive),
                reactive.texto,
                row["selected_label"],
                reactive.evidencia_sugerida,
                _control_names(row) if uses_document_controls else _document_names(row),
                row["observacion"],
                (
                    control_files or "Sin archivos vinculados"
                    if uses_document_controls
                    else ", ".join(e.archivo_nombre_original for e in row["evidence"]) or "Sin archivos"
                ),
            ]
        )
    _style_table(findings, findings_header)
    _set_widths(
        findings,
        {"A": 24, "B": 11, "C": 13, "D": 12, "E": 58, "F": 15, "G": 42, "H": 38, "I": 40, "J": 28},
    )

    guide = workbook.create_sheet("Guía")
    _write_title(guide, "Guía de lectura ISO 45001", _evaluation_scope_description(evaluation), 4)
    guide.append([])
    guide.append(["Aviso", "Diagnóstico de preparación. No constituye certificación ni sustituye la norma autorizada o la legislación aplicable."])
    guide.append([])
    guide.append(["Rango", "Nivel de madurez", "Interpretación"])
    guide_header = guide.max_row
    for interval, level in MATURITY_GUIDE:
        guide.append([interval, level, "La madurez se calcula con respuestas No, Parcial y Sí sobre todos los reactivos."])
    guide.append([])
    guide.append(["Cláusula", "Enfoque de evaluación", "% Cumplimiento", "Lectura"])
    clause_header = guide.max_row
    for clause in summary["clauses"]:
        guide.append(
            [
                clause["numero"],
                CLAUSE_GUIDANCE.get(clause["numero"], clause["nombre"]),
                _excel_percent(clause["percent"]),
                clause["maturity_label"],
            ]
        )
    _style_table(guide, guide_header)
    _style_table(guide, clause_header)
    _set_widths(guide, {"A": 16, "B": 46, "C": 16, "D": 34})

    if uses_document_controls:
        traceability = workbook.create_sheet("Trazabilidad documental")
        _write_title(
            traceability,
            "Trazabilidad archivo - control documental - reactivos",
            _evaluation_scope_description(evaluation),
            8,
        )
        traceability.append([])
        traceability.append(
            [
                "Archivo",
                "Control",
                "Tipo",
                "Estado",
                "Puntos cubiertos",
                "Reactivos cubiertos",
                "Observación",
                "Contenido mínimo",
            ]
        )
        traceability_header = traceability.max_row
        for control in summary["document_controls"]:
            files = control["evidencias"] or [None]
            for evidence in files:
                traceability.append(
                    [
                        evidence["archivo_nombre_original"] if evidence else "Sin archivo vinculado",
                        f"{control['codigo']} - {control['nombre']}",
                        control["clasificacion"],
                        control["estado_label"],
                        f"{control['puntos_cubiertos']}/{control['total_puntos']}",
                        ", ".join(item["codigo"] for item in control["reactivos"]) or "-",
                        control["observacion"] or "",
                        control["contenido_minimo"],
                    ]
                )
        _style_table(traceability, traceability_header)
        _set_widths(traceability, {"A": 30, "B": 42, "C": 20, "D": 18, "E": 18, "F": 48, "G": 44, "H": 54})

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A4"
        sheet.sheet_view.showGridLines = False

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_iso45001_pdf(evaluation) -> BytesIO:
    """Build the executive ISO 45001 report while preserving the XLSX contract."""
    summary = summarize_iso45001_evaluation(evaluation)
    uses_document_controls = _uses_document_control_coverage(summary)
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.55 * inch,
        rightMargin=0.55 * inch,
        topMargin=0.60 * inch,
        bottomMargin=0.55 * inch,
    )
    styles = _build_iso45001_pdf_styles()
    story: list = []
    story.extend(_build_iso45001_cover_story(evaluation, summary, styles))
    story.extend(_build_iso45001_index_story(evaluation, summary, styles))

    story.append(Paragraph("Resumen ejecutivo", styles["section"]))
    if not summary["is_final"]:
        story.append(
            Paragraph(
                "Resultado preliminar. La interpretaci\u00f3n puede cambiar mientras la captura o la revisi\u00f3n permanezcan abiertas.",
                styles["warning"],
            )
        )
        story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_metric_cards(summary, styles))
    story.append(Spacer(1, 0.14 * inch))
    story.append(_build_iso45001_methodology_panel(summary, styles))
    story.append(Spacer(1, 0.12 * inch))
    story.append(Paragraph("Lectura ejecutiva", styles["subsection"]))
    story.append(Paragraph(_paragraph_escape(_global_maturity_comment(summary)), styles["body"]))

    story.append(CondPageBreak(4.3 * inch))
    story.append(Spacer(1, 0.12 * inch))
    story.append(Paragraph("Gr\u00e1ficas de autodiagn\u00f3stico", styles["section"]))
    story.append(_build_iso45001_chart_pair(summary))
    story.append(Spacer(1, 0.10 * inch))
    story.append(_build_iso45001_progress_chart(summary))
    story.append(Spacer(1, 0.12 * inch))
    story.append(_build_iso45001_radar_chart(summary))

    story.append(CondPageBreak(4.4 * inch))
    story.append(Paragraph("Diagn\u00f3stico accionable", styles["section"]))
    story.append(
        Paragraph(
            "La prioridad combina cumplimiento, avance de captura, brechas derivadas y soporte de evidencia exigido. "
            "Los resultados no sustituyen la evaluaci\u00f3n de riesgos ni la verificaci\u00f3n legal aplicable.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_clause_heatmap(summary, styles))
    story.append(Spacer(1, 0.12 * inch))
    story.append(Paragraph("Apartados de atenci\u00f3n prioritaria", styles["subsection"]))
    story.append(_build_iso45001_priority_sections_table(_priority_sections(summary), styles))
    story.append(Spacer(1, 0.12 * inch))
    story.append(
        Paragraph(
            "Cobertura documental por cl\u00e1usula" if uses_document_controls else "Cobertura de adjuntos exigidos por cl\u00e1usula",
            styles["subsection"],
        )
    )
    story.append(_build_iso45001_evidence_coverage_table(summary, styles))

    story.append(CondPageBreak(4.2 * inch))
    story.append(Paragraph("Hallazgos prioritarios", styles["section"]))
    story.append(
        Paragraph(
            "Brecha prioritaria: respuesta No. Brecha de mejora: respuesta Parcial. Esta selecci\u00f3n muestra las diez primeras "
            "brechas ordenadas para revisi\u00f3n ejecutiva; el anexo conserva el detalle completo.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_priority_findings_table(summary, styles))

    story.append(CondPageBreak(3.9 * inch))
    story.append(Paragraph("Amd. 1:2024 - acci\u00f3n clim\u00e1tica", styles["section"]))
    story.append(
        Paragraph(
            "La enmienda incorpora la consideraci\u00f3n de cambio clim\u00e1tico en el contexto y en los requisitos de partes interesadas. "
            "Los dos controles se muestran por separado para facilitar su seguimiento.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_climate_panel(summary, styles))

    story.append(CondPageBreak(4.2 * inch))
    story.append(Paragraph("Información documentada y soporte", styles["section"]))
    story.append(
        Paragraph(
            (
                "La cobertura se evalúa una vez por control: 31 controles normativos y seis grupos complementarios. "
                "Cada archivo puede vincularse a varios controles y reactivos, sin duplicar cargas ni alterar el puntaje ISO."
                if uses_document_controls
                else "El catálogo es una referencia normativa y de trazabilidad. El estado presentado se deriva de los reactivos vinculados; "
                "no crea una evaluación documental independiente."
            ),
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_document_metrics(summary, styles))
    story.append(Spacer(1, 0.10 * inch))
    story.append(Paragraph("Documentos que requieren atenci\u00f3n", styles["subsection"]))
    story.append(_build_iso45001_document_attention_table(summary, styles))

    story.append(CondPageBreak(4.4 * inch))
    story.append(Paragraph("Gu\u00eda de madurez por cl\u00e1usula", styles["section"]))
    story.append(
        Paragraph(
            "Las interpretaciones traducen las respuestas capturadas en se\u00f1ales de madurez y siguientes pasos de SST. "
            "Son orientativas y deben contrastarse con evidencia objetiva y el contexto operativo.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_clause_guidance_table(summary, styles))

    for clause in summary["clauses"]:
        story.append(CondPageBreak(3.9 * inch))
        story.extend(_build_iso45001_clause_story(clause, styles))

    story.append(PageBreak())
    story.append(Paragraph("Anexo A - Cat\u00e1logo de informaci\u00f3n documentada", styles["section"]))
    story.append(
        Paragraph(
            (
                "Trazabilidad completa entre cada control, sus puntos mínimos, archivos reutilizables y reactivos vinculados."
                if uses_document_controls
                else "Trazabilidad completa entre cada elemento documental, reactivos vinculados y adjuntos exigidos/satisfechos."
            ),
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_document_appendix(summary, styles))

    story.append(PageBreak())
    story.append(Paragraph("Anexo B - Hallazgos y oportunidades de mejora", styles["section"]))
    story.append(
        Paragraph(
            "Relaci\u00f3n completa de respuestas No y Parcial, con observaciones, evidencia sugerida y documentos relacionados.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_findings_appendix(summary, styles))

    if uses_document_controls:
        story.append(PageBreak())
        story.append(Paragraph("Anexo C - Trazabilidad documental", styles["section"]))
        story.append(
            Paragraph(
                "Relación archivo -> control documental -> puntos cubiertos -> reactivos sustentados.",
                styles["body"],
            )
        )
        story.append(Spacer(1, 0.08 * inch))
        story.append(_build_iso45001_document_traceability_appendix(summary, styles))

    story.append(PageBreak())
    story.append(
        Paragraph(
            "Anexo D - Trazabilidad de reactivos" if uses_document_controls else "Anexo C - Trazabilidad de reactivos",
            styles["section"],
        )
    )
    story.append(
        Paragraph(
            (
                "Matriz consultable de los 308 reactivos, respuesta, observaci\u00f3n y controles documentales vinculados."
                if uses_document_controls
                else "Matriz consultable de los 308 reactivos, respuesta, observaci\u00f3n y archivos adjuntos disponibles."
            ),
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_build_iso45001_traceability_appendix(summary, styles))

    document.build(story, onFirstPage=_draw_iso45001_pdf_chrome, onLaterPages=_draw_iso45001_pdf_chrome)
    buffer.seek(0)
    return buffer


def _legacy_build_iso45001_pdf(evaluation) -> BytesIO:
    summary = summarize_iso45001_evaluation(evaluation)
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.55 * inch,
    )
    styles = _pdf_styles()
    story = []
    story.extend(_cover_story(evaluation, summary, styles))
    story.append(Paragraph("Resumen ejecutivo", styles["section"]))
    story.append(_metric_table(summary, styles))
    story.append(Spacer(1, 0.15 * inch))
    if not summary["is_final"]:
        story.append(
            Paragraph(
                "Resultado preliminar. La captura o revisión puede cambiar antes del cierre oficial.",
                styles["warning"],
            )
        )
        story.append(Spacer(1, 0.1 * inch))
    story.append(Paragraph("Cumplimiento por cláusula", styles["subsection"]))
    story.append(_clause_chart(summary))
    story.append(Spacer(1, 0.12 * inch))
    story.append(_clause_table(summary, styles))
    story.append(Spacer(1, 0.15 * inch))
    story.append(Paragraph("Matriz de madurez", styles["subsection"]))
    story.append(_radar_chart(summary))
    story.append(
        Paragraph(
            "La matriz presenta el porcentaje de cumplimiento por cláusula. Las respuestas No y Parcial alimentan los hallazgos derivados.",
            styles["body"],
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("Información documentada requerida", styles["section"]))
    story.append(
        Paragraph(
            "La siguiente relación es normativa y sirve para identificar evidencia. No representa un expediente documental ni una puntuación independiente.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_document_table(summary, styles))

    story.append(PageBreak())
    story.append(Paragraph("Hallazgos y oportunidades de mejora", styles["section"]))
    story.append(
        Paragraph(
            "Brecha prioritaria: respuesta No. Brecha de mejora: respuesta Parcial. Los archivos adjuntos se muestran como evidencia disponible.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.08 * inch))
    story.append(_findings_table(summary, styles))

    story.append(PageBreak())
    story.append(Paragraph("Guía y anexo de trazabilidad", styles["section"]))
    story.append(
        Paragraph(
            "Este informe es un diagnóstico de preparación para ISO 45001:2018 y su Amd. 1:2024. Para una auditoría formal debe consultarse la norma autorizada y la legislación aplicable.",
            styles["note"],
        )
    )
    story.append(Spacer(1, 0.1 * inch))
    story.append(_guide_table(summary, styles))
    story.append(Spacer(1, 0.14 * inch))
    story.append(Paragraph("Anexo de reactivos", styles["subsection"]))
    story.append(_appendix_table(summary, styles))
    document.build(story, onFirstPage=_draw_pdf_chrome, onLaterPages=_draw_pdf_chrome)
    buffer.seek(0)
    return buffer


def _write_title(sheet, title: str, subtitle: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=16, bold=True, color=EXCEL_COLORS["white"])
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=EXCEL_COLORS["primary"])
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    sheet.cell(2, 1, subtitle)
    sheet.cell(2, 1).font = Font(size=10, italic=True, color=EXCEL_COLORS["primary"])
    sheet.cell(2, 1).alignment = Alignment(horizontal="center")


def _style_table(sheet, header_row: int) -> None:
    header_fill = PatternFill("solid", fgColor=EXCEL_COLORS["primary"])
    border = Border(
        left=Side(style="thin", color=EXCEL_COLORS["line"]),
        right=Side(style="thin", color=EXCEL_COLORS["line"]),
        top=Side(style="thin", color=EXCEL_COLORS["line"]),
        bottom=Side(style="thin", color=EXCEL_COLORS["line"]),
    )
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color=EXCEL_COLORS["white"])
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"


def _set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _reactive_identifier(reactive) -> str:
    return getattr(reactive, "identificador", None) or getattr(reactive, "codigo", None) or f"R-{reactive.numero:03d}"


def _documents_for_row(row: dict) -> list:
    return row.get("documents") or list(getattr(row["reactivo"], "documentos_requeridos", []) or [])


def _document_names(row: dict) -> str:
    documents = _documents_for_row(row)
    return "; ".join(f"{document.codigo} - {document.nombre}" for document in documents) or "No requiere"


def _uses_document_control_coverage(summary: dict) -> bool:
    return bool(summary.get("uses_document_control_coverage"))


def _control_names(row: dict) -> str:
    controls = list(getattr(row["reactivo"], "controles_evidencia", []) or [])
    return "; ".join(f"{control.codigo} - {control.nombre}" for control in controls) or "Evidencia complementaria"


def _document_rows(summary: dict) -> list[tuple[object, list[dict]]]:
    version = summary["evaluation"].ciclo.version
    catalog = list(getattr(version, "documentos_obligatorios", []) or [])
    rows_by_id = {document.id: [] for document in catalog}
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                for document in _documents_for_row(row):
                    rows_by_id.setdefault(document.id, []).append(row)
    return [(document, rows_by_id.get(document.id, [])) for document in catalog]


def _document_observation(rows: list[dict]) -> str:
    if not rows:
        return "Sin reactivos vinculados."
    if any(row["selected"] == "no" for row in rows):
        return "Hay brechas prioritarias relacionadas."
    if any(row["selected"] == "parcial" for row in rows):
        return "Hay evidencia o implementación parcial."
    if any(row["selected"] == "si" and not row["evidence"] for row in rows):
        return "Respuesta positiva sin archivo adjunto."
    return "Con evidencia disponible o sin respuesta positiva."


def _findings(summary: dict) -> list[tuple[dict, dict, dict]]:
    rows = []
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                if row["selected"] in {"no", "parcial"}:
                    rows.append((clause, section, row))
    return sorted(rows, key=lambda item: (0 if item[2]["selected"] == "no" else 1, item[0]["numero"], item[1]["codigo"], item[2]["reactivo"].numero))


def _finding_label(row: dict) -> str:
    if row["selected"] == "no":
        return "Brecha prioritaria"
    if row["selected"] == "parcial":
        return "Brecha de mejora"
    if not row["answered"]:
        return "Sin evaluar"
    return "Conforme"


def _percent_display(value) -> str:
    return "-" if value is None else f"{value:.2f}%"


def _excel_percent(value):
    return "" if value is None else value / 100


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_kicker": ParagraphStyle("Iso45001CoverKicker", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=10, leading=13, alignment=TA_CENTER, textColor=PDF_THEME["muted"]),
        "cover_title": ParagraphStyle("Iso45001CoverTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=25, leading=30, alignment=TA_CENTER, textColor=PDF_THEME["primary"], spaceAfter=8),
        "cover_subject": ParagraphStyle("Iso45001CoverSubject", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=15, leading=18, alignment=TA_CENTER, textColor=PDF_THEME["ink"], spaceAfter=6),
        "cover_meta": ParagraphStyle("Iso45001CoverMeta", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_CENTER, textColor=PDF_THEME["muted"]),
        "section": ParagraphStyle("Iso45001Section", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=PDF_THEME["primary"], spaceBefore=3, spaceAfter=8),
        "subsection": ParagraphStyle("Iso45001Subsection", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=PDF_THEME["ink"], spaceBefore=4, spaceAfter=5),
        "body": ParagraphStyle("Iso45001Body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.2, leading=13, textColor=PDF_THEME["ink"]),
        "small": ParagraphStyle("Iso45001Small", parent=base["BodyText"], fontName="Helvetica", fontSize=7.3, leading=9.2, textColor=PDF_THEME["ink"]),
        "note": ParagraphStyle("Iso45001Note", parent=base["BodyText"], fontName="Helvetica", fontSize=8.4, leading=11.5, textColor=PDF_THEME["muted"], backColor=PDF_THEME["primary_soft"], borderColor=PDF_THEME["line"], borderWidth=0.5, borderPadding=7),
        "warning": ParagraphStyle("Iso45001Warning", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=8.5, leading=11.5, textColor=PDF_THEME["red"], backColor=colors.HexColor("#FCE4E4"), borderColor=PDF_THEME["red"], borderWidth=0.5, borderPadding=7),
        "table_head": ParagraphStyle("Iso45001TableHead", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=7.2, leading=8.5, textColor=colors.white),
    }


def _cover_story(evaluation, summary: dict, styles: dict) -> list:
    return [
        Spacer(1, 0.5 * inch),
        Paragraph("SISTEMA DE GESTIÓN DE SEGURIDAD Y SALUD EN EL TRABAJO", styles["cover_kicker"]),
        Paragraph("Diagnóstico ISO 45001:2018", styles["cover_title"]),
        Paragraph("incluye Amd. 1:2024 - acción climática", styles["cover_kicker"]),
        Spacer(1, 0.14 * inch),
        Paragraph(escape(_evaluation_scope_name(evaluation)), styles["cover_subject"]),
        Paragraph(escape(_evaluation_scope_description(evaluation)), styles["cover_meta"]),
        Paragraph(escape(evaluation.ciclo.nombre), styles["cover_meta"]),
        Spacer(1, 0.26 * inch),
        Paragraph(
            f"Estado: <b>{escape(summary['state_label'])}</b><br/>"
            f"Avance: <b>{_percent_display(summary['completion'])}</b><br/>"
            f"Cumplimiento: <b>{_percent_display(summary['percent'])}</b><br/>"
            f"Madurez: <b>{escape(summary['maturity_label'])}</b><br/>"
            f"Generado: {escape(format_iso45001_datetime(evaluation.updated_at))}",
            styles["cover_meta"],
        ),
        Spacer(1, 0.45 * inch),
        Paragraph(
            "Herramienta de diagnóstico y preparación. No constituye una certificación ni sustituye la norma autorizada o la legislación aplicable.",
            styles["note"],
        ),
        PageBreak(),
    ]


def _metric_table(summary: dict, styles: dict) -> Table:
    rows = [
        ["Avance", _percent_display(summary["completion"])],
        ["Cumplimiento", _percent_display(summary["percent"])],
        ["Madurez", summary["maturity_label"]],
        ["Reactivos", f"{summary['answered_questions']} de {summary['total_questions']}"],
        ["Evidencias", str(summary["evidence_count"])],
        ["Cobertura requerida", _percent_display(summary["evidence_coverage"]["percent"])],
    ]
    table = Table(rows, colWidths=[2.1 * inch, 4.25 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), PDF_THEME["primary_soft"]),
                ("TEXTCOLOR", (0, 0), (-1, -1), PDF_THEME["ink"]),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _clause_chart(summary: dict) -> Drawing:
    drawing = Drawing(500, 165)
    x0, y0, width, height = 58, 32, 404, 108
    drawing.add(Line(x0, y0, x0 + width, y0, strokeColor=PDF_THEME["line"]))
    for index, clause in enumerate(summary["clauses"]):
        x = x0 + index * (width / max(len(summary["clauses"]), 1)) + 14
        bar_width = 30
        percent = clause["percent"] or 0
        bar_height = height * (percent / 100)
        color = _percent_color(percent)
        drawing.add(Rect(x, y0, bar_width, bar_height, fillColor=color, strokeColor=color))
        drawing.add(String(x + bar_width / 2, y0 - 15, clause["numero"], textAnchor="middle", fontName="Helvetica-Bold", fontSize=8, fillColor=PDF_THEME["ink"]))
        drawing.add(String(x + bar_width / 2, y0 + bar_height + 5, f"{percent:.0f}%", textAnchor="middle", fontName="Helvetica", fontSize=7, fillColor=PDF_THEME["muted"]))
    drawing.add(String(x0, y0 + height + 15, "Porcentaje de cumplimiento por cláusula", fontName="Helvetica-Bold", fontSize=9, fillColor=PDF_THEME["ink"]))
    return drawing


def _clause_table(summary: dict, styles: dict) -> Table:
    data = [[
        Paragraph("Cláusula", styles["table_head"]),
        Paragraph("Reactivos", styles["table_head"]),
        Paragraph("Respondidos", styles["table_head"]),
        Paragraph("Cumplimiento", styles["table_head"]),
        Paragraph("Madurez", styles["table_head"]),
        Paragraph("Evidencias", styles["table_head"]),
    ]]
    for clause in summary["clauses"]:
        data.append(
            [
                Paragraph(f"<b>{escape(clause['numero'])}</b><br/>{escape(clause['nombre'])}", styles["small"]),
                str(clause["total"]),
                str(clause["answered"]),
                _percent_display(clause["percent"]),
                clause["maturity_label"],
                str(clause.get("evidence_count", 0)),
            ]
        )
    table = Table(data, colWidths=[2.4 * inch, 0.62 * inch, 0.75 * inch, 0.85 * inch, 1.25 * inch, 0.7 * inch], repeatRows=1)
    _style_pdf_table(table)
    return table


def _radar_chart(summary: dict) -> Drawing:
    drawing = Drawing(500, 245)
    center_x, center_y, radius = 250, 115, 78
    clauses = summary["clauses"]
    count = max(len(clauses), 1)
    for factor in (0.25, 0.5, 0.75, 1):
        points = []
        for index in range(count):
            angle = math.pi / 2 + (2 * math.pi * index / count)
            points.extend([center_x + radius * factor * math.cos(angle), center_y + radius * factor * math.sin(angle)])
        drawing.add(Polygon(points, strokeColor=PDF_THEME["line"], fillColor=None, strokeWidth=0.4))
    actual = []
    for index, clause in enumerate(clauses):
        angle = math.pi / 2 + (2 * math.pi * index / count)
        x = center_x + radius * math.cos(angle)
        y = center_y + radius * math.sin(angle)
        drawing.add(Line(center_x, center_y, x, y, strokeColor=PDF_THEME["line"], strokeWidth=0.5))
        label_x = center_x + (radius + 22) * math.cos(angle)
        label_y = center_y + (radius + 22) * math.sin(angle)
        drawing.add(String(label_x, label_y, clause["numero"], textAnchor="middle", fontName="Helvetica-Bold", fontSize=8, fillColor=PDF_THEME["ink"]))
        percent = clause["percent"] or 0
        actual.extend([center_x + radius * (percent / 100) * math.cos(angle), center_y + radius * (percent / 100) * math.sin(angle)])
    if actual:
        drawing.add(Polygon(actual, strokeColor=PDF_THEME["primary"], fillColor=colors.Color(0.12, 0.44, 0.47, alpha=0.22), strokeWidth=1.6))
    drawing.add(String(250, 223, "Radar de cumplimiento por cláusula", textAnchor="middle", fontName="Helvetica-Bold", fontSize=10, fillColor=PDF_THEME["ink"]))
    return drawing


def _document_table(summary: dict, styles: dict) -> Table:
    data = [[
        Paragraph("Código", styles["table_head"]),
        Paragraph("Apartado", styles["table_head"]),
        Paragraph("Información documentada", styles["table_head"]),
        Paragraph("Cobertura", styles["table_head"]),
        Paragraph("Hallazgo", styles["table_head"]),
    ]]
    for document, rows in _document_rows(summary):
        selected = [row for row in rows if row["selected"] in {"si", "parcial"}]
        evidence = [row for row in selected if row["evidence"]]
        coverage = len(evidence) / len(selected) * 100 if selected else None
        data.append(
            [
                document.codigo,
                document.apartado_codigo,
                Paragraph(f"<b>{escape(document.nombre)}</b><br/>{escape(document.tipo)}", styles["small"]),
                _percent_display(coverage),
                Paragraph(escape(_document_observation(rows)), styles["small"]),
            ]
        )
    table = Table(data, colWidths=[0.52 * inch, 0.64 * inch, 2.85 * inch, 0.7 * inch, 2.0 * inch], repeatRows=1)
    _style_pdf_table(table)
    return table


def _findings_table(summary: dict, styles: dict) -> Table:
    findings = _findings(summary)
    data = [[
        Paragraph("Prioridad", styles["table_head"]),
        Paragraph("Ref.", styles["table_head"]),
        Paragraph("Hallazgo", styles["table_head"]),
        Paragraph("Documentos / evidencia", styles["table_head"]),
    ]]
    for clause, section, row in findings:
        reactive = row["reactivo"]
        supporting = _document_names(row)
        files = ", ".join(e.archivo_nombre_original for e in row["evidence"]) or "Sin archivos"
        data.append(
            [
                Paragraph(escape(_finding_label(row)), styles["small"]),
                Paragraph(f"{escape(section['codigo'])}<br/>{escape(_reactive_identifier(reactive))}", styles["small"]),
                Paragraph(escape(reactive.texto), styles["small"]),
                Paragraph(f"<b>Documentos:</b> {escape(supporting)}<br/><b>Archivos:</b> {escape(files)}", styles["small"]),
            ]
        )
    if len(data) == 1:
        data.append(["Sin hallazgos", "-", "Todos los reactivos respondidos como Sí.", "-"])
    table = Table(data, colWidths=[1.05 * inch, 0.68 * inch, 2.8 * inch, 2.1 * inch], repeatRows=1)
    _style_pdf_table(table)
    return table


def _guide_table(summary: dict, styles: dict) -> Table:
    data = [[
        Paragraph("Rango", styles["table_head"]),
        Paragraph("Madurez", styles["table_head"]),
        Paragraph("Cláusula", styles["table_head"]),
        Paragraph("Enfoque", styles["table_head"]),
    ]]
    for index, clause in enumerate(summary["clauses"]):
        interval, level = MATURITY_GUIDE[min(index, len(MATURITY_GUIDE) - 1)]
        data.append([interval, level, clause["numero"], Paragraph(escape(CLAUSE_GUIDANCE.get(clause["numero"], clause["nombre"])), styles["small"])])
    table = Table(data, colWidths=[0.75 * inch, 1.55 * inch, 0.65 * inch, 3.7 * inch], repeatRows=1)
    _style_pdf_table(table)
    return table


def _appendix_table(summary: dict, styles: dict) -> Table:
    data = [[
        Paragraph("Ref.", styles["table_head"]),
        Paragraph("Reactivo", styles["table_head"]),
        Paragraph("Respuesta", styles["table_head"]),
        Paragraph("Archivos", styles["table_head"]),
    ]]
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                reactive = row["reactivo"]
                data.append(
                    [
                        f"{section['codigo']} / {_reactive_identifier(reactive)}",
                        Paragraph(escape(reactive.texto), styles["small"]),
                        row["selected_label"],
                        Paragraph(escape(", ".join(e.archivo_nombre_original for e in row["evidence"]) or "-"), styles["small"]),
                    ]
                )
    table = Table(data, colWidths=[0.9 * inch, 3.8 * inch, 0.75 * inch, 1.2 * inch], repeatRows=1)
    _style_pdf_table(table)
    return table


def _style_pdf_table(table: Table) -> None:
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_THEME["primary"]),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FBFB")]),
            ]
        )
    )


def _percent_color(percent: float) -> colors.Color:
    if percent < 50:
        return PDF_THEME["red"]
    if percent < 80:
        return PDF_THEME["yellow"]
    return PDF_THEME["green"]


def _draw_pdf_chrome(canvas, document) -> None:
    canvas.saveState()
    canvas.setStrokeColor(PDF_THEME["line"])
    canvas.line(document.leftMargin, 0.42 * inch, letter[0] - document.rightMargin, 0.42 * inch)
    canvas.setFont("Helvetica", 7.3)
    canvas.setFillColor(PDF_THEME["muted"])
    canvas.drawString(document.leftMargin, 0.27 * inch, "Diagnóstico ISO 45001:2018 + Amd. 1:2024")
    canvas.drawRightString(letter[0] - document.rightMargin, 0.27 * inch, f"Página {canvas.getPageNumber()}")
    canvas.restoreState()


# Executive PDF helpers. They remain separate from the legacy helpers above
# so the spreadsheet export keeps its current behavior.
def _register_iso45001_pdf_fonts() -> tuple[str, str]:
    package_root = Path(__file__).resolve().parents[1]
    regular_env = os.environ.get("ISO45001_PDF_FONT_REGULAR")
    bold_env = os.environ.get("ISO45001_PDF_FONT_BOLD")
    regular_candidates = [
        Path(regular_env) if regular_env else None,
        package_root / "static" / "fonts" / "GoogleSans-Regular.ttf",
        package_root / "static" / "fonts" / "ProductSans-Regular.ttf",
        Path("C:/Windows/Fonts/GoogleSans-Regular.ttf"),
        Path("C:/Windows/Fonts/ProductSans-Regular.ttf"),
        Path("/usr/share/fonts/truetype/google-sans/GoogleSans-Regular.ttf"),
        Path("/usr/share/fonts/truetype/product-sans/ProductSans-Regular.ttf"),
    ]
    bold_candidates = [
        Path(bold_env) if bold_env else None,
        package_root / "static" / "fonts" / "GoogleSans-Bold.ttf",
        package_root / "static" / "fonts" / "ProductSans-Bold.ttf",
        Path("C:/Windows/Fonts/GoogleSans-Bold.ttf"),
        Path("C:/Windows/Fonts/ProductSans-Bold.ttf"),
        Path("/usr/share/fonts/truetype/google-sans/GoogleSans-Bold.ttf"),
        Path("/usr/share/fonts/truetype/product-sans/ProductSans-Bold.ttf"),
    ]
    regular_path = next((path for path in regular_candidates if path and path.exists()), None)
    bold_path = next((path for path in bold_candidates if path and path.exists()), regular_path)
    if regular_path:
        try:
            pdfmetrics.registerFont(TTFont("Iso45001Sans", str(regular_path)))
            pdfmetrics.registerFont(TTFont("Iso45001Sans-Bold", str(bold_path or regular_path)))
            return "Iso45001Sans", "Iso45001Sans-Bold"
        except Exception:
            pass
    return "Helvetica", "Helvetica-Bold"


ISO45001_PDF_FONT_REGULAR, ISO45001_PDF_FONT_BOLD = _register_iso45001_pdf_fonts()


def _build_iso45001_pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_kicker": ParagraphStyle(
            "Iso45001ExecutiveCoverKicker",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=9.8,
            leading=12,
            textColor=PDF_THEME["muted"],
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "cover_title": ParagraphStyle(
            "Iso45001ExecutiveCoverTitle",
            parent=base["Title"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=30,
            leading=35,
            textColor=PDF_THEME["primary"],
            alignment=TA_CENTER,
            spaceAfter=9,
        ),
        "cover_subject": ParagraphStyle(
            "Iso45001ExecutiveCoverSubject",
            parent=base["Heading2"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=16,
            leading=20,
            textColor=PDF_THEME["ink"],
            alignment=TA_CENTER,
            spaceAfter=15,
        ),
        "cover_org": ParagraphStyle(
            "Iso45001ExecutiveCoverOrg",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=12,
            leading=15,
            textColor=PDF_THEME["ink"],
            alignment=TA_CENTER,
        ),
        "cover_department": ParagraphStyle(
            "Iso45001ExecutiveCoverDepartment",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=9.4,
            leading=12,
            textColor=PDF_THEME["primary"],
            alignment=TA_CENTER,
        ),
        "cover_note": ParagraphStyle(
            "Iso45001ExecutiveCoverNote",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=10,
            leading=13,
            textColor=PDF_THEME["muted"],
            alignment=TA_CENTER,
        ),
        "cover_meta": ParagraphStyle(
            "Iso45001ExecutiveCoverMeta",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=8.2,
            leading=10.2,
            textColor=PDF_THEME["muted"],
            alignment=TA_CENTER,
        ),
        "index_lead": ParagraphStyle(
            "Iso45001ExecutiveIndexLead",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=9.1,
            leading=12.1,
            textColor=PDF_THEME["muted"],
            spaceAfter=6,
        ),
        "section": ParagraphStyle(
            "Iso45001ExecutiveSection",
            parent=base["Heading2"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=14,
            leading=18,
            textColor=PDF_THEME["primary"],
            spaceBefore=4,
            spaceAfter=6,
        ),
        "subsection": ParagraphStyle(
            "Iso45001ExecutiveSubsection",
            parent=base["Heading3"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=11.1,
            leading=14,
            textColor=PDF_THEME["primary_dark"],
            spaceBefore=2,
            spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "Iso45001ExecutiveBody",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=9,
            leading=12,
            textColor=PDF_THEME["ink"],
        ),
        "note": ParagraphStyle(
            "Iso45001ExecutiveNote",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=7.7,
            leading=9.5,
            textColor=PDF_THEME["muted"],
        ),
        "small": ParagraphStyle(
            "Iso45001ExecutiveSmall",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=7.65,
            leading=9.35,
            textColor=PDF_THEME["ink"],
        ),
        "small_center": ParagraphStyle(
            "Iso45001ExecutiveSmallCenter",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=7.65,
            leading=9.35,
            textColor=PDF_THEME["ink"],
            alignment=TA_CENTER,
        ),
        "table_header": ParagraphStyle(
            "Iso45001ExecutiveTableHeader",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=7.25,
            leading=8.8,
            textColor=colors.white,
            alignment=TA_CENTER,
        ),
        "metric_label": ParagraphStyle(
            "Iso45001ExecutiveMetricLabel",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=7.2,
            leading=8.8,
            textColor=PDF_THEME["muted"],
            alignment=TA_CENTER,
        ),
        "metric_value": ParagraphStyle(
            "Iso45001ExecutiveMetricValue",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=17,
            leading=19,
            textColor=PDF_THEME["primary_dark"],
            alignment=TA_CENTER,
        ),
        "metric_caption": ParagraphStyle(
            "Iso45001ExecutiveMetricCaption",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_REGULAR,
            fontSize=6.9,
            leading=8.2,
            textColor=PDF_THEME["muted"],
            alignment=TA_CENTER,
        ),
        "warning": ParagraphStyle(
            "Iso45001ExecutiveWarning",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=8.4,
            leading=11,
            textColor=PDF_THEME["red"],
            backColor=PDF_THEME["red_soft"],
            borderColor=PDF_THEME["red"],
            borderWidth=0.45,
            borderPadding=7,
        ),
        "traffic_light": ParagraphStyle(
            "Iso45001ExecutiveTrafficLight",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=7.3,
            leading=8.8,
            textColor=colors.white,
            alignment=TA_CENTER,
        ),
        "traffic_dark": ParagraphStyle(
            "Iso45001ExecutiveTrafficDark",
            parent=base["BodyText"],
            fontName=ISO45001_PDF_FONT_BOLD,
            fontSize=7.3,
            leading=8.8,
            textColor=PDF_THEME["ink"],
            alignment=TA_CENTER,
        ),
    }


def _build_iso45001_cover_story(evaluation, summary: dict, styles: dict[str, ParagraphStyle]) -> list:
    status = "OFICIAL" if summary["is_final"] else "PRELIMINAR"
    status_color = PDF_THEME["primary_dark"] if summary["is_final"] else PDF_THEME["yellow"]
    institutional_header = Table(
        [
            [Paragraph("H. Ayuntamiento de Hermosillo", styles["cover_org"])],
            [Paragraph("Direcci\u00f3n de Recursos Humanos", styles["cover_department"])],
        ],
        colWidths=[7.05 * inch],
        rowHeights=[0.28 * inch, 0.24 * inch],
    )
    institutional_header.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_THEME["primary_soft"]),
                ("LINEBELOW", (0, 1), (-1, 1), 1.6, PDF_THEME["primary"]),
                ("BOX", (0, 0), (-1, -1), 0.45, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    status_badge = Table(
        [[Paragraph(status, styles["table_header"])]],
        colWidths=[1.55 * inch],
        rowHeights=[0.32 * inch],
    )
    status_badge.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), status_color),
                ("BOX", (0, 0), (-1, -1), 0.4, status_color),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    return [
        Spacer(1, 0.42 * inch),
        institutional_header,
        Spacer(1, 1.13 * inch),
        Paragraph("Sistema de gesti\u00f3n de seguridad y salud en el trabajo", styles["cover_kicker"]),
        Paragraph("Autodiagn\u00f3stico ISO 45001:2018", styles["cover_title"]),
        Paragraph("incluye Amd. 1:2024 - acci\u00f3n clim\u00e1tica", styles["cover_kicker"]),
        Paragraph(_paragraph_escape(_evaluation_scope_name(evaluation)), styles["cover_subject"]),
        Paragraph(_paragraph_escape(_evaluation_scope_description(evaluation)), styles["cover_meta"]),
        status_badge,
        Spacer(1, 0.38 * inch),
        Paragraph(
            "Informe ejecutivo de implementaci\u00f3n, evidencia documental, participaci\u00f3n de trabajadores y madurez del SGSST.",
            styles["cover_note"],
        ),
        Spacer(1, 0.22 * inch),
        Paragraph(
            f"Ciclo: {_paragraph_escape(evaluation.ciclo.nombre)} | Estado: {_paragraph_escape(summary['state_label'])}",
            styles["cover_meta"],
        ),
        Spacer(1, 2.12 * inch),
        Paragraph(
            "Diagn\u00f3stico de preparaci\u00f3n. No constituye certificaci\u00f3n ni sustituye la norma autorizada o la legislaci\u00f3n aplicable.",
            styles["cover_note"],
        ),
        PageBreak(),
    ]


def _build_iso45001_index_story(evaluation, summary: dict, styles: dict[str, ParagraphStyle]) -> list:
    rows = [
        ["Orden", "Secci\u00f3n", "Qu\u00e9 consultar"],
        ["1", "Resumen ejecutivo", "Indicadores, m\u00e9todo de c\u00e1lculo y lectura global."],
        ["2", "Gr\u00e1ficas", "Cumplimiento por cl\u00e1usula, respuestas, avance y radar de madurez."],
        ["3", "Diagn\u00f3stico accionable", "Matriz de madurez, apartados prioritarios y adjuntos exigidos."],
        ["4", "Hallazgos prioritarios", "Brechas principales, evidencia sugerida y documentos relacionados."],
        ["5", "Amd. 1:2024", "Controles clim\u00e1ticos incorporados en las cl\u00e1usulas 4.1 y 4.2."],
        ["6", "Gu\u00eda y detalle", "Interpretaci\u00f3n por cl\u00e1usula y siguiente paso de SST."],
        ["7", "Anexos", "Cat\u00e1logo documental, hallazgos completos y trazabilidad por reactivo."],
    ]
    return [
        Paragraph("\u00cdndice del reporte", styles["section"]),
        Paragraph(
            "Este documento est\u00e1 organizado para revisi\u00f3n ejecutiva, seguimiento operativo y trazabilidad de la evidencia del SGSST.",
            styles["index_lead"],
        ),
        _build_iso45001_metadata_panel(evaluation, summary, styles),
        Spacer(1, 0.15 * inch),
        _iso45001_pdf_table(rows, styles, col_widths=[0.62 * inch, 1.72 * inch, 4.71 * inch]),
        Spacer(1, 0.14 * inch),
        Paragraph(
            "\u00cdndice de madurez: cumplimiento = puntos obtenidos / (308 reactivos x 2). "
            "No=0, Parcial=1 y S\u00ed=2; no se admite N/A en este diagn\u00f3stico.",
            styles["note"],
        ),
        PageBreak(),
    ]


def _build_iso45001_metadata_panel(evaluation, summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    responsible = evaluation.responsable.nombre if getattr(evaluation, "responsable", None) else "Sin responsable"
    reviewer = evaluation.revisor.nombre if getattr(evaluation, "revisor", None) else "Sin revisor"
    cycle_dates = f"{_iso45001_date_or_dash(evaluation.ciclo.fecha_inicio)} al {_iso45001_date_or_dash(evaluation.ciclo.fecha_cierre)}"
    rows = [
        ["Unidad administrativa", "Dependencia", "Ciclo", "Estado"],
        [
            _evaluation_unit_name(evaluation),
            _evaluation_dependency_name(evaluation),
            evaluation.ciclo.nombre,
            summary["state_label"],
        ],
        ["Periodo", "Responsable", "Revisor", "\u00cdndice de madurez"],
        [
            cycle_dates,
            responsible,
            reviewer,
            f"{_format_iso45001_percent(summary['percent'])} - {summary['maturity_label']}",
        ],
        ["Captura", "Versi\u00f3n", "Reactivos", "Descarga"],
        [
            _format_iso45001_percent(summary["completion"]),
            evaluation.ciclo.version.nombre,
            str(summary["total_questions"]),
            format_iso45001_datetime(utcnow()),
        ],
    ]
    table_rows = []
    for row_index, row in enumerate(rows):
        style = styles["table_header"] if row_index % 2 == 0 else styles["small_center"]
        table_rows.append([_iso45001_as_paragraph(value, style) for value in row])
    table = Table(table_rows, colWidths=[2.15 * inch, 2.05 * inch, 1.65 * inch, 1.2 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_THEME["primary"]),
                ("BACKGROUND", (0, 2), (-1, 2), PDF_THEME["primary"]),
                ("BACKGROUND", (0, 4), (-1, 4), PDF_THEME["primary"]),
                ("BACKGROUND", (0, 1), (-1, 1), PDF_THEME["primary_soft"]),
                ("BACKGROUND", (0, 3), (-1, 3), colors.white),
                ("BACKGROUND", (0, 5), (-1, 5), PDF_THEME["primary_soft"]),
                ("GRID", (0, 0), (-1, -1), 0.45, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _iso45001_executive_extension_marker() -> None:
    """Anchor for the professional PDF helpers below."""


def _paragraph_escape(value) -> str:
    return escape(str(value or ""), quote=False)


def _iso45001_document_status(document: dict) -> dict:
    questions = document.get("questions", [])
    selected = [row["selected"] for row in questions]
    no_count = sum(answer == "no" for answer in selected)
    partial_count = sum(answer == "parcial" for answer in selected)
    answered_count = sum(answer is not None for answer in selected)
    required = document.get("evidence_required", 0)
    satisfied = document.get("evidence_satisfied", 0)
    missing = max(required - satisfied, 0)
    attachment_ratio = f"{satisfied}/{required}" if required else "No exigidos"
    if no_count:
        return {
            "document": document,
            "kind": "priority",
            "priority_order": 0,
            "label": "Brecha prioritaria",
            "attachment_ratio": attachment_ratio,
            "reading": f"{no_count} respuesta(s) No vinculada(s); revisar el soporte normativo y operativo.",
        }
    if partial_count:
        return {
            "document": document,
            "kind": "partial",
            "priority_order": 1,
            "label": "Implementaci\u00f3n parcial",
            "attachment_ratio": attachment_ratio,
            "reading": f"{partial_count} respuesta(s) Parcial vinculada(s); completar la implementaci\u00f3n y su evidencia.",
        }
    if missing:
        return {
            "document": document,
            "kind": "pending",
            "priority_order": 2,
            "label": "Adjunto pendiente",
            "attachment_ratio": attachment_ratio,
            "reading": f"Faltan {missing} adjunto(s) exigido(s) para respuestas positivas vinculadas.",
        }
    if questions and answered_count < len(questions):
        return {
            "document": document,
            "kind": "unassessed",
            "priority_order": 3,
            "label": "Sin evaluar",
            "attachment_ratio": attachment_ratio,
            "reading": "Hay reactivos vinculados sin respuesta; complete la captura antes de concluir.",
        }
    if questions:
        return {
            "document": document,
            "kind": "supported",
            "priority_order": 4,
            "label": "Sustentado",
            "attachment_ratio": attachment_ratio,
            "reading": "Sin respuesta No o Parcial vinculada y con el soporte exigido disponible.",
        }
    return {
        "document": document,
        "kind": "unassessed",
        "priority_order": 3,
        "label": "Sin reactivos vinculados",
        "attachment_ratio": attachment_ratio,
        "reading": "Elemento del cat\u00e1logo sin trazabilidad a un reactivo del diagn\u00f3stico.",
    }


def _draw_iso45001_pdf_chrome(canvas, document) -> None:
    canvas.saveState()
    width, height = document.pagesize
    canvas.setFillColor(PDF_THEME["primary"])
    canvas.rect(document.leftMargin, height - 0.32 * inch, width - document.leftMargin - document.rightMargin, 0.10 * inch, fill=1, stroke=0)
    canvas.setStrokeColor(PDF_THEME["line"])
    canvas.line(document.leftMargin, 0.42 * inch, width - document.rightMargin, 0.42 * inch)
    canvas.setFillColor(PDF_THEME["muted"])
    canvas.setFont(ISO45001_PDF_FONT_REGULAR, 7.2)
    canvas.drawString(document.leftMargin, 0.28 * inch, ISO45001_PDF_FOOTER)
    canvas.drawRightString(width - document.rightMargin, 0.28 * inch, f"P\u00e1gina {canvas.getPageNumber()}")
    canvas.setTitle("Diagn\u00f3stico ISO 45001:2018 + Amd. 1:2024")
    canvas.setSubject("Informe ejecutivo de preparaci\u00f3n del sistema de gesti\u00f3n de SST")
    canvas.restoreState()


def _iso45001_maturity_profile(percent: float | None) -> dict[str, str]:
    if percent is None:
        return {
            "reading": "Sin base de c\u00e1lculo para interpretar la madurez.",
            "action": "Complete la captura y valide el alcance del diagn\u00f3stico.",
        }
    if percent == 0:
        return {
            "reading": "Madurez no iniciada: no se observa implementaci\u00f3n acreditada en los reactivos evaluados.",
            "action": "Defina responsables, controles m\u00ednimos, registros y una fecha de arranque verificable.",
        }
    if percent <= 20:
        return {
            "reading": "Madurez inicial: existen se\u00f1ales aisladas, pero la gesti\u00f3n depende de pr\u00e1cticas reactivas o informales.",
            "action": "Formalice el m\u00e9todo, asigne responsables y conserve evidencia objetiva suficiente.",
        }
    if percent <= 40:
        return {
            "reading": "Madurez en desarrollo: hay pr\u00e1cticas parciales, pero falta consistencia, despliegue o trazabilidad.",
            "action": "Convierta las pr\u00e1cticas parciales en controles, registros, indicadores y revisiones peri\u00f3dicas.",
        }
    if percent <= 60:
        return {
            "reading": "Madurez definida: el requisito opera parcialmente y requiere fortalecer evidencia, eficacia y seguimiento.",
            "action": "Cierre brechas de implementaci\u00f3n, verifique eficacia y conecte los controles con riesgos y objetivos de SST.",
        }
    if percent <= 80:
        return {
            "reading": "Madurez gestionada: el requisito funciona con evidencia razonable, aunque puede reforzarse la medici\u00f3n.",
            "action": "Use auditor\u00edas, an\u00e1lisis de datos y revisi\u00f3n directiva para sostener el desempe\u00f1o.",
        }
    return {
        "reading": "Madurez optimizada: el requisito est\u00e1 integrado al trabajo y cuenta con soporte consistente.",
        "action": "Mantenga el seguimiento, comparta buenas pr\u00e1cticas y documente mejoras preventivas.",
    }


def _global_maturity_comment(summary: dict) -> str:
    profile = _iso45001_maturity_profile(summary["percent"])
    counts = summary["response_counts"]
    if _uses_document_control_coverage(summary):
        stats = summary["document_evidence_stats"]
        return (
            f"{profile['reading']} El corte registra {counts['no']} respuestas No, {counts['parcial']} Parcial, "
            f"{counts['si']} Sí y {counts['sin_respuesta']} sin respuesta. "
            f"La cobertura documental registra {stats['covered_points']} de {stats['total_points']} puntos, "
            f"{stats['sustained_controls']} controles sustentados y {stats['file_count']} archivo(s) único(s). "
            f"{profile['action']}"
        )
    evidence = summary["evidence_coverage"]
    return (
        f"{profile['reading']} El corte registra {counts['no']} respuestas No, {counts['parcial']} Parcial, "
        f"{counts['si']} S\u00ed y {counts['sin_respuesta']} sin respuesta. "
        f"Los adjuntos exigidos est\u00e1n cubiertos en {evidence['satisfied']} de {evidence['required']} casos. "
        f"{profile['action']}"
    )


def _iso45001_clause_guidance_row(clause: dict) -> dict:
    guide = CLAUSE_REPORT_GUIDANCE.get(
        str(clause["numero"]),
        {
            "title": clause["nombre"],
            "focus": "los requisitos del SGSST",
            "support": "Revise el requisito mediante evidencia objetiva y trazable.",
            "evidence": "Registros y documentos vigentes relacionados con el alcance declarado.",
        },
    )
    profile = _iso45001_maturity_profile(clause["percent"])
    counts = _iso45001_clause_response_counts(clause)
    coverage = clause["evidence_coverage"]
    document_coverage = clause.get("document_control_coverage")
    support_text = (
        (
            f"{document_coverage['file_count']} archivo(s) documental(es) \u00fanico(s) vinculado(s)."
            if document_coverage["total_controls"]
            else "Sin controles documentales independientes."
        )
        if document_coverage is not None
        else f"{clause['evidence_count']} archivo(s)."
    )
    reading = (
        f"{profile['reading']} El foco es {guide['focus']}. "
        f"Corte: {counts['no']} No, {counts['parcial']} Parcial, {counts['si']} S\u00ed, "
        f"{counts['sin_respuesta']} sin respuesta y {support_text}"
    )
    if document_coverage is not None:
        document_reading = (
            "Sin controles documentales independientes en esta cl\u00e1usula."
            if not document_coverage["total_controls"]
            else (
                f"Cobertura documental: {document_coverage['sustained_controls']}/"
                f"{document_coverage['total_controls']} controles sustentados; "
                f"{document_coverage['covered_points']}/{document_coverage['total_points']} puntos cubiertos."
            )
        )
        next_step = f"{profile['action']} {document_reading} Evidencia orientativa: {guide['evidence']}"
    else:
        next_step = (
            f"{profile['action']} Adjuntos exigidos: {coverage['satisfied']}/{coverage['required']}. "
            f"Evidencia orientativa: {guide['evidence']}"
        )
    return {
        "clause": f"{clause['numero']} {clause['nombre']}",
        "maturity": clause["maturity_label"],
        "reading": _clip_iso45001(reading, 455),
        "next_step": _clip_iso45001(next_step, 335),
    }


def _iso45001_clause_support_row(clause: dict) -> dict:
    guide = CLAUSE_REPORT_GUIDANCE.get(
        str(clause["numero"]),
        {
            "title": clause["nombre"],
            "support": "Revise el requisito mediante evidencia objetiva y trazable.",
            "evidence": "Registros y documentos vigentes relacionados con el alcance declarado.",
        },
    )
    return {"title": guide["title"], "support": guide["support"], "evidence": guide["evidence"]}


def _iso45001_section_maturity_comment(section: dict) -> str:
    guide = CLAUSE_REPORT_GUIDANCE.get(str(section["clausula"]), {"focus": "los requisitos del SGSST"})
    profile = _iso45001_maturity_profile(section["percent"])
    counts = _iso45001_section_response_counts(section)
    coverage = section["evidence_coverage"]
    document_coverage = section.get("document_control_coverage")
    capture = "Captura completa." if section["completion"] >= 100 else f"Complete la captura ({_format_iso45001_percent(section['completion'])})."
    if document_coverage is not None:
        document_text = (
            "No tiene un control documental independiente en esta matriz."
            if not document_coverage["total_controls"]
            else (
                f"Cobertura documental: {document_coverage['sustained_controls']}/"
                f"{document_coverage['total_controls']} controles sustentados y "
                f"{document_coverage['covered_points']}/{document_coverage['total_points']} puntos cubiertos."
            )
        )
    else:
        document_text = f"Adjuntos exigidos: {coverage['satisfied']}/{coverage['required']}."
    return _clip_iso45001(
        f"{profile['reading']} {capture} Para este apartado revise {guide['focus']}; "
        f"hay {counts['no']} No, {counts['parcial']} Parcial y {counts['sin_respuesta']} sin respuesta. "
        f"{document_text} Siguiente paso: {profile['action']}",
        520,
    )


def _iso45001_finding_records(summary: dict) -> list[dict]:
    records = []
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                finding = row.get("finding")
                if finding is None:
                    continue
                record = dict(finding)
                record["row"] = row
                record["evidence"] = row["evidence"]
                records.append(record)
    records.sort(
        key=lambda item: (
            item["priority_order"],
            str(item["clausula"]),
            str(item["apartado"]),
            item["reactivo"].orden,
        )
    )
    return records


def _iso45001_support_text(finding: dict) -> str:
    return _iso45001_support_text_for_controls(
        finding,
        document_controls=None,
    )


def _iso45001_support_text_for_controls(
    finding: dict,
    document_controls: list[dict] | None,
) -> str:
    if document_controls is not None:
        controls = _document_controls_for_reactive(
            document_controls,
            finding["reactivo"].id,
        )
        control_text = (
            "; ".join(f"{control['codigo']} - {control['nombre']}" for control in controls)
            or "Sin control documental trazado"
        )
        files = _document_control_file_names(controls) or "Sin archivos vinculados"
        return (
            f"<b>Controles:</b> {_clip_iso45001(control_text, 155)}"
            f"<br/><b>Archivos vinculados:</b> {_clip_iso45001(files, 120)}"
        )
    documents = finding.get("documents") or []
    document_text = "; ".join(f"{document.codigo} - {document.nombre}" for document in documents) or "No requiere"
    files = ", ".join(evidence.archivo_nombre_original for evidence in finding.get("evidence", [])) or "Sin archivos"
    return f"<b>Documentos:</b> {_clip_iso45001(document_text, 155)}<br/><b>Archivos:</b> {_clip_iso45001(files, 120)}"


def _document_controls_for_reactive(
    document_controls: list[dict],
    reactive_id: int,
) -> list[dict]:
    return [
        control
        for control in document_controls
        if any(item["id"] == reactive_id for item in control.get("reactivos", []))
    ]


def _document_control_file_names(controls: list[dict]) -> str:
    names: list[str] = []
    for control in controls:
        for evidence in control.get("evidencias", []):
            name = evidence["archivo_nombre_original"]
            if name not in names:
                names.append(name)
    return ", ".join(names)


def _iso45001_climate_rows(summary: dict) -> list[dict]:
    rows = []
    for section in summary["sections"]:
        for row in section["questions"]:
            if getattr(row["reactivo"], "es_enmienda_2024", False):
                rows.append(row)
    return sorted(rows, key=lambda row: row["reactivo"].orden)


def _iso45001_climate_interpretation(row: dict) -> str:
    if row["selected"] == "si":
        if row["evidence"]:
            return "Control incorporado y con evidencia adjunta. Mantenga su revisi\u00f3n ante cambios relevantes."
        return "Control declarado como implementado; valide que la evidencia objetiva permanezca disponible."
    if row["selected"] == "parcial":
        return "Implementaci\u00f3n parcial. Formalice el an\u00e1lisis, responsables y la evidencia de revisi\u00f3n."
    if row["selected"] == "no":
        return "Brecha prioritaria de la enmienda. Integre el an\u00e1lisis clim\u00e1tico al contexto o a las partes interesadas."
    return "Pendiente de captura. No hay base suficiente para confirmar la consideraci\u00f3n de la enmienda."


def _iso45001_pdf_table(
    rows: list[list],
    styles: dict[str, ParagraphStyle],
    *,
    col_widths: list | None = None,
    repeat_rows: int = 1,
    header: bool = True,
    center_from_col: int | None = None,
) -> Table:
    table_rows = []
    for row_index, row in enumerate(rows):
        style = styles["table_header"] if header and row_index == 0 else styles["small"]
        table_rows.append([_iso45001_as_paragraph(value, style) for value in row])
    table = Table(table_rows, colWidths=col_widths, repeatRows=repeat_rows if header else 0, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.4, PDF_THEME["line"]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [colors.white, PDF_THEME["paper"]]),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_THEME["primary"]),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
            ]
        )
    if center_from_col is not None:
        commands.append(("ALIGN", (center_from_col, 1 if header else 0), (-1, -1), "CENTER"))
    table.setStyle(TableStyle(commands))
    return table


def _iso45001_as_paragraph(value, style: ParagraphStyle):
    if isinstance(value, Paragraph):
        return value
    text = _paragraph_escape(value)
    # The report builders use these three tags for deliberate, local layout;
    # all other potentially unsafe markup remains escaped.
    text = (
        text.replace("&lt;br/&gt;", "<br/>")
        .replace("&lt;b&gt;", "<b>")
        .replace("&lt;/b&gt;", "</b>")
    )
    return Paragraph(text, style)


def _format_iso45001_percent(value: float | int | None) -> str:
    if value is None:
        return "-"
    numeric = round(float(value), 2)
    return f"{numeric:.2f}".rstrip("0").rstrip(".") + "%"


def _iso45001_date_or_dash(value) -> str:
    return value.strftime("%d/%m/%Y") if value else "-"


def _clip_iso45001(value, limit: int) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return f"{text[: max(limit - 3, 0)].rstrip()}..."


def _compact_iso45001_maturity(label: str | None) -> str:
    text = str(label or "-")
    return text.split(" - ", 1)[1] if " - " in text else text


def _maturity_level_prefix_iso45001(label: str | None) -> str:
    text = str(label or "")
    return text.split(" - ", 1)[0] if " - " in text else "Nivel global"


def _iso45001_traffic_color(percent: float | int | None):
    if percent is None:
        return PDF_THEME["gray"]
    value = float(percent)
    if value <= 40:
        return PDF_THEME["red"]
    if value <= 80:
        return PDF_THEME["yellow"]
    return PDF_THEME["green"]


def _iso45001_traffic_soft_color(percent: float | int | None):
    if percent is None:
        return PDF_THEME["gray_soft"]
    value = float(percent)
    if value <= 40:
        return PDF_THEME["red_soft"]
    if value <= 80:
        return PDF_THEME["yellow_soft"]
    return PDF_THEME["green_soft"]


def _iso45001_maturity_bucket(percent: float | None) -> int | None:
    if percent is None:
        return None
    if percent == 0:
        return 0
    if percent <= 20:
        return 1
    if percent <= 40:
        return 2
    if percent <= 60:
        return 3
    if percent <= 80:
        return 4
    return 5


def _iso45001_response_distribution(summary: dict) -> list[dict]:
    counts = summary["response_counts"]
    total = sum(counts.get(key, 0) for key, _label, _color in RESPONSE_ORDER)
    rows = []
    for key, label, color in RESPONSE_ORDER:
        count = counts.get(key, 0)
        percent = round((count / total) * 100, 2) if total else 0.0
        rows.append({"key": key, "label": label, "color": color, "count": count, "percent": percent})
    return rows


def _iso45001_section_response_counts(section: dict) -> dict[str, int]:
    counts = {"no": 0, "parcial": 0, "si": 0, "sin_respuesta": 0}
    for row in section["questions"]:
        selected = row["selected"] or "sin_respuesta"
        if selected in counts:
            counts[selected] += 1
    return counts


def _iso45001_clause_response_counts(clause: dict) -> dict[str, int]:
    counts = {"no": 0, "parcial": 0, "si": 0, "sin_respuesta": 0}
    for section in clause["sections"]:
        section_counts = _iso45001_section_response_counts(section)
        for key in counts:
            counts[key] += section_counts[key]
    return counts


def _priority_sections(summary: dict, limit: int = 10) -> list[dict]:
    sections = [section for section in summary["sections"] if section["percent"] is not None]
    sections.sort(key=lambda section: (section["percent"], section["completion"], section["codigo"]))
    return sections[:limit]


def _iso45001_priority_label(percent: float | None, completion: float | None) -> str:
    if completion is not None and completion < 100:
        return "Completar captura"
    if percent is None:
        return "Sin base"
    if percent <= 40:
        return "Alta"
    if percent <= 70:
        return "Media"
    return "Seguimiento"


def _iso45001_evidence_reading(coverage: dict) -> str:
    if coverage["required"] == 0:
        return "Sin adjuntos exigidos"
    if coverage["missing"] == 0:
        return "Soporte exigido completo"
    return f"Faltan {coverage['missing']} adjunto(s) exigido(s)"


def _build_iso45001_clause_guidance_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Cl\u00e1usula", "Madurez", "Lectura", "Siguiente paso"]]
    for clause in summary["clauses"]:
        guidance = _iso45001_clause_guidance_row(clause)
        rows.append([guidance["clause"], guidance["maturity"], guidance["reading"], guidance["next_step"]])
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[1.08 * inch, 1.14 * inch, 2.92 * inch, 2.11 * inch],
    )


def _build_iso45001_clause_story(clause: dict, styles: dict[str, ParagraphStyle]) -> list:
    guidance = _iso45001_clause_guidance_row(clause)
    support = _iso45001_clause_support_row(clause)
    attention = [
        section
        for section in clause["sections"]
        if section["percent"] is not None and (section["percent"] < 80 or section["completion"] < 100)
    ]
    attention.sort(key=lambda item: (item["percent"], item["completion"], item["codigo"]))
    story = [
        Paragraph(f"Cl\u00e1usula {clause['numero']} - {_paragraph_escape(clause['nombre'])}", styles["section"]),
        _build_iso45001_clause_metrics_table(clause, styles),
        Spacer(1, 0.10 * inch),
        Paragraph("Interpretaci\u00f3n", styles["subsection"]),
        Paragraph(
            f"{_paragraph_escape(guidance['reading'])} <b>Siguiente paso:</b> {_paragraph_escape(guidance['next_step'])}",
            styles["body"],
        ),
        Spacer(1, 0.08 * inch),
        Paragraph("Soporte normativo de referencia", styles["subsection"]),
        Paragraph(
            f"<b>{_paragraph_escape(support['title'])}:</b> {_paragraph_escape(support['support'])} "
            f"<b>Evidencia orientativa:</b> {_paragraph_escape(support['evidence'])}",
            styles["body"],
        ),
        Spacer(1, 0.10 * inch),
        Paragraph("Apartados evaluados", styles["subsection"]),
        _build_iso45001_clause_sections_table(clause, styles),
        Spacer(1, 0.10 * inch),
        Paragraph("\u00c1reas de atenci\u00f3n", styles["subsection"]),
    ]
    if attention:
        for section in attention[:5]:
            story.append(
                Paragraph(
                    f"- {_paragraph_escape(section['codigo'])} {_paragraph_escape(section['nombre'])}: "
                    f"{_paragraph_escape(_iso45001_section_maturity_comment(section))}",
                    styles["body"],
                )
            )
    else:
        story.append(Paragraph("Sin apartados cr\u00edticos con la informaci\u00f3n capturada.", styles["body"]))
    return story


def _build_iso45001_clause_metrics_table(clause: dict, styles: dict[str, ParagraphStyle]) -> Table:
    coverage = clause["evidence_coverage"]
    document_coverage = clause.get("document_control_coverage")
    document_label = "Doc. control" if document_coverage is not None else "Adjuntos"
    document_value = (
        (
            f"{document_coverage['sustained_controls']}/"
            f"{document_coverage['total_controls']}"
            if document_coverage["total_controls"]
            else "-"
        )
        if document_coverage is not None
        else f"{coverage['satisfied']}/{coverage['required']}"
    )
    rows = [
        ["Reactivos", "Respondidos", "Puntos", "Cumplimiento", document_label, "Madurez"],
        [
            str(clause["total"]),
            str(clause["answered"]),
            str(clause["points"]),
            _format_iso45001_percent(clause["percent"]),
            document_value,
            clause["maturity_label"],
        ],
    ]
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.85 * inch, 0.92 * inch, 0.62 * inch, 1.05 * inch, 0.95 * inch, 2.66 * inch],
        center_from_col=0,
    )


def _build_iso45001_clause_sections_table(clause: dict, styles: dict[str, ParagraphStyle]) -> Table:
    uses_document_controls = any(
        section.get("document_control_coverage") is not None
        for section in clause["sections"]
    )
    support_header = "Doc." if uses_document_controls else "Adj."
    rows = [["Apartado", "React.", "Avance", "Cumpl.", "No", "Parcial", support_header, "Madurez"]]
    for section in clause["sections"]:
        counts = _iso45001_section_response_counts(section)
        coverage = section["evidence_coverage"]
        document_coverage = section.get("document_control_coverage")
        support_value = (
            (
                f"{document_coverage['sustained_controls']}/"
                f"{document_coverage['total_controls']}"
                if document_coverage["total_controls"]
                else "-"
            )
            if document_coverage is not None
            else f"{coverage['satisfied']}/{coverage['required']}"
        )
        rows.append(
            [
                f"{section['codigo']} {section['nombre']}",
                str(section["total"]),
                _format_iso45001_percent(section["completion"]),
                _format_iso45001_percent(section["percent"]),
                str(counts["no"]),
                str(counts["parcial"]),
                support_value,
                section["maturity_label"],
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[2.31 * inch, 0.46 * inch, 0.68 * inch, 0.68 * inch, 0.42 * inch, 0.57 * inch, 0.58 * inch, 1.57 * inch],
        center_from_col=1,
    )


def _build_iso45001_document_appendix(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        return _build_iso45001_document_control_appendix(summary, styles)
    rows = [["C\u00f3d.", "Apartado", "Documento / tipo", "React.", "Adjuntos", "Arch.", "Estado", "Contenido m\u00ednimo"]]
    for document in summary["document_requirements"]:
        status = _iso45001_document_status(document)
        rows.append(
            [
                document["codigo"],
                document["apartado"],
                f"{document['nombre']}<br/>{document['clasificacion']}",
                str(document["mapped_reactives"]),
                status["attachment_ratio"],
                str(document["evidence_count"]),
                status["label"],
                document["contenido_minimo"],
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.45 * inch, 0.55 * inch, 1.88 * inch, 0.42 * inch, 0.67 * inch, 0.4 * inch, 1.05 * inch, 1.83 * inch],
        center_from_col=3,
    )


def _build_iso45001_findings_appendix(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    uses_document_controls = _uses_document_control_coverage(summary)
    document_controls = summary.get("document_controls") if uses_document_controls else None
    rows = [[
        "Prioridad",
        "Ref.",
        "Hallazgo / observaci\u00f3n",
        "Evidencia sugerida",
        "Controles / archivos" if uses_document_controls else "Documentos / archivos",
    ]]
    findings = _iso45001_finding_records(summary)
    if not findings:
        rows.append(["Sin hallazgos", "-", "No se registran respuestas No o Parcial.", "-", "-"])
    for finding in findings:
        reactive = finding["reactivo"]
        observation = _clip_iso45001(finding["observacion"] or "Sin observaci\u00f3n capturada.", 200)
        rows.append(
            [
                finding["label"],
                f"{finding['apartado']}<br/>{_reactive_identifier(reactive)}",
                f"{reactive.texto}<br/><b>Observaci\u00f3n:</b> {observation}",
                _clip_iso45001(finding["evidencia_sugerida"], 260),
                _iso45001_support_text_for_controls(finding, document_controls),
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.92 * inch, 0.68 * inch, 2.15 * inch, 1.85 * inch, 1.58 * inch],
        center_from_col=1,
    )


def _build_iso45001_traceability_appendix(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    uses_document_controls = _uses_document_control_coverage(summary)
    document_controls = summary.get("document_controls") if uses_document_controls else None
    rows = [[
        "Ref.",
        "Reactivo",
        "Respuesta",
        "Observaci\u00f3n / controles" if uses_document_controls else "Observaci\u00f3n / archivos",
    ]]
    for clause in summary["clauses"]:
        for section in clause["sections"]:
            for row in section["questions"]:
                reactive = row["reactivo"]
                observation = _clip_iso45001(row["observacion"] or "Sin observaci\u00f3n", 180)
                if uses_document_controls:
                    support = _iso45001_support_text_for_controls(
                        {"reactivo": reactive, "evidence": row["evidence"]},
                        document_controls,
                    )
                else:
                    files = ", ".join(evidence.archivo_nombre_original for evidence in row["evidence"]) or "Sin archivos"
                    support = f"<b>Archivos:</b> {_clip_iso45001(files, 120)}"
                rows.append(
                    [
                        f"{section['codigo']}<br/>{_reactive_identifier(reactive)}",
                        reactive.texto,
                        row["selected_label"],
                        f"<b>Observaci\u00f3n:</b> {observation}<br/>{support}",
                    ]
                )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.82 * inch, 3.55 * inch, 0.72 * inch, 2.09 * inch],
        center_from_col=2,
    )


def _build_iso45001_priority_findings_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Prioridad", "Ref.", "Hallazgo / observaci\u00f3n", "Evidencia sugerida", "Soporte"]]
    findings = _iso45001_finding_records(summary)
    if not findings:
        rows.append(["Sin hallazgos", "-", "No se registran respuestas No o Parcial.", "-", "-"])
    for finding in findings[:10]:
        reactive = finding["reactivo"]
        observation = _clip_iso45001(finding["observacion"] or "Sin observaci\u00f3n capturada.", 160)
        finding_text = f"{reactive.texto}<br/><b>Observaci\u00f3n:</b> {observation}"
        support = _iso45001_support_text_for_controls(
            finding,
            summary.get("document_controls") if _uses_document_control_coverage(summary) else None,
        )
        rows.append(
            [
                finding["label"],
                f"{finding['apartado']}<br/>{_reactive_identifier(reactive)}",
                finding_text,
                _clip_iso45001(finding["evidencia_sugerida"], 220),
                support,
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.95 * inch, 0.68 * inch, 2.35 * inch, 2.0 * inch, 1.25 * inch],
        center_from_col=1,
    )


def _build_iso45001_climate_panel(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Control Amd. 1:2024", "Resultado", "Soporte disponible", "Interpretaci\u00f3n"]]
    climate_rows = _iso45001_climate_rows(summary)
    if not climate_rows:
        rows.append(["Controles clim\u00e1ticos no disponibles en el cat\u00e1logo.", "-", "-", "Valide la versi\u00f3n del diagn\u00f3stico."])
    for row in climate_rows:
        reactive = row["reactivo"]
        files = ", ".join(evidence.archivo_nombre_original for evidence in row["evidence"]) or "Sin archivos"
        rows.append(
            [
                f"{row['apartado']} / {_reactive_identifier(reactive)}<br/>{reactive.texto}",
                row["selected_label"],
                f"{len(row['evidence'])} archivo(s)<br/>{_clip_iso45001(files, 120)}",
                _iso45001_climate_interpretation(row),
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[2.55 * inch, 0.82 * inch, 1.35 * inch, 2.53 * inch],
        center_from_col=1,
    )


def _build_iso45001_document_metrics(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        return _build_iso45001_document_control_metric_cards(summary, styles)
    statuses = [_iso45001_document_status(document) for document in summary["document_requirements"]]
    priority = sum(1 for item in statuses if item["kind"] == "priority")
    partial = sum(1 for item in statuses if item["kind"] == "partial")
    supported = sum(1 for item in statuses if item["kind"] == "supported")
    pending = sum(1 for item in statuses if item["kind"] in {"pending", "unassessed"})
    evidence = summary["evidence_coverage"]
    cards = [
        ("Cat\u00e1logo", str(len(statuses)), "Elementos de referencia"),
        ("Brecha prioritaria", str(priority), "Documentos con al menos un No"),
        ("Implementaci\u00f3n parcial", str(partial), "Documentos vinculados a Parcial"),
        ("Sustentados", str(supported), "Sin brecha y con soporte requerido"),
        ("Pendientes", str(pending), "Sin respuesta o adjunto exigido"),
        ("Adjuntos exigidos", f"{evidence['satisfied']}/{evidence['required']}", "Cobertura de soporte"),
    ]
    cells = []
    for label, value, caption in cards:
        card = Table(
            [
                [Paragraph(label.upper(), styles["metric_label"])],
                [Paragraph(value, styles["metric_value"])],
                [Paragraph(caption, styles["metric_caption"])],
            ],
            colWidths=[1.12 * inch],
            rowHeights=[0.22 * inch, 0.42 * inch, 0.27 * inch],
        )
        card.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        cells.append(card)
    table = Table([cells], colWidths=[1.16 * inch] * 6, rowHeights=[0.98 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_THEME["primary_soft"]),
                ("BOX", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("INNERGRID", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return table


def _build_iso45001_document_attention_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        return _build_iso45001_document_control_attention_table(summary, styles)
    statuses = [_iso45001_document_status(document) for document in summary["document_requirements"]]
    attention = [item for item in statuses if item["kind"] != "supported"]
    attention.sort(key=lambda item: (item["priority_order"], item["document"]["orden"], item["document"]["codigo"]))
    rows = [["C\u00f3digo", "Documento / referencia", "React.", "Adjuntos exigidos", "Estado", "Lectura"]]
    if not attention:
        rows.append(["-", "Sin documentos con brecha o pendiente.", "-", "-", "Sustentado", "El soporte requerido est\u00e1 disponible."])
    for item in attention[:10]:
        document = item["document"]
        rows.append(
            [
                document["codigo"],
                f"{document['apartado']} {document['nombre']}<br/>{document['clasificacion']}",
                str(document["mapped_reactives"]),
                item["attachment_ratio"],
                item["label"],
                item["reading"],
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.55 * inch, 2.35 * inch, 0.48 * inch, 0.95 * inch, 1.1 * inch, 1.62 * inch],
        center_from_col=2,
    )


def _document_controls_for_clause(summary: dict, clause_number: str) -> list[dict]:
    return [
        control
        for control in summary.get("document_controls", [])
        if str(control.get("clausula")) == str(clause_number)
    ]


def _document_control_scope_stats(controls: list[dict]) -> dict:
    total_points = sum(int(control.get("total_puntos") or 0) for control in controls)
    covered_points = sum(int(control.get("puntos_cubiertos") or 0) for control in controls)
    evidence_ids = {
        evidence["id"]
        for control in controls
        for evidence in control.get("evidencias", [])
    }
    evaluated = sum(bool(control.get("evaluated")) for control in controls)
    sustained = sum(
        1
        for control in controls
        if control.get("estado") in {"parcial", "si"} and control.get("evidence_ids")
    )
    return {
        "total": len(controls),
        "evaluated": evaluated,
        "covered_points": covered_points,
        "total_points": total_points,
        "percent": round((covered_points / total_points) * 100, 2) if total_points else 0.0,
        "files": len(evidence_ids),
        "sustained": sustained,
        "gaps": sum(control.get("estado") == "no" for control in controls),
    }


def _document_control_status(control: dict) -> tuple[int, str, str]:
    if not control.get("evaluated"):
        return 0, "Pendiente de evaluar", "Complete los puntos de cobertura o documente la brecha."
    if control.get("estado") == "no":
        return 1, "Brecha documental", control.get("observacion") or "Registre la brecha documental."
    if not control.get("evidence_ids"):
        return 2, "Archivo pendiente", "La cobertura declarada requiere al menos un archivo vinculado."
    if control.get("estado") == "parcial":
        return 3, "Cobertura parcial", "Complete los puntos mínimos pendientes y mantenga la evidencia vinculada."
    return 4, "Sustentado", "Todos los puntos están declarados cubiertos y cuentan con archivo vinculado."


def _build_iso45001_v2_metric_cards(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    """Executive cards that keep ISO scoring separate from document coverage."""

    counts = summary["response_counts"]
    stats = summary["document_evidence_stats"]
    cards = [
        ("Avance", _format_iso45001_percent(summary["completion"]), f"{summary['answered_questions']} de {summary['total_questions']} reactivos"),
        ("Cumplimiento", _format_iso45001_percent(summary["percent"]), "Puntos sobre el máximo posible"),
        ("Madurez", _compact_iso45001_maturity(summary["maturity_label"]), _maturity_level_prefix_iso45001(summary["maturity_label"])),
        ("Brechas No", str(counts["no"]), "Atención prioritaria"),
        ("Controles", f"{stats['evaluated_controls']}/{stats['total_controls']}", "Documentales evaluados"),
        ("Puntos doc.", f"{stats['covered_points']}/{stats['total_points']}", "Cobertura declarada"),
        ("Sustentados", str(stats["sustained_controls"]), "Con archivo vinculado"),
        ("Archivos", str(stats["file_count"]), "Únicos y reutilizables"),
    ]
    return _build_iso45001_cards_grid(cards, styles)


def _build_iso45001_document_control_metric_cards(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    stats = summary["document_evidence_stats"]
    cards = [
        ("Controles", f"{stats['evaluated_controls']}/{stats['total_controls']}", "Evaluados"),
        ("Puntos", f"{stats['covered_points']}/{stats['total_points']}", "Cobertura declarada"),
        ("Sustentados", str(stats["sustained_controls"]), "Con archivo vinculado"),
        ("Archivos", str(stats["file_count"]), "Únicos y reutilizables"),
        ("Brechas", str(stats["gap_controls"]), "Controles sin puntos"),
        ("Parcial", str(stats["partial_controls"]), "Cobertura por completar"),
        ("Normativos", str(stats["normative_controls"]), "Documentos D-01 a D-31"),
        ("Complementarios", str(stats["complementary_controls"]), "Grupos G-01 a G-06"),
    ]
    return _build_iso45001_cards_grid(cards, styles)


def _build_iso45001_cards_grid(
    cards: list[tuple[str, str, str]],
    styles: dict[str, ParagraphStyle],
) -> Table:
    rows = []
    for start in range(0, len(cards), 4):
        row = []
        for label, value, caption in cards[start : start + 4]:
            card = Table(
                [
                    [Paragraph(_paragraph_escape(label.upper()), styles["metric_label"])],
                    [Paragraph(_paragraph_escape(value), styles["metric_value"])],
                    [Paragraph(_paragraph_escape(caption), styles["metric_caption"])],
                ],
                colWidths=[1.55 * inch],
                rowHeights=[0.22 * inch, 0.42 * inch, 0.28 * inch],
            )
            card.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
            row.append(card)
        rows.append(row)
    table = Table(rows, colWidths=[1.76 * inch] * 4, rowHeights=[0.98 * inch] * len(rows))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, PDF_THEME["primary_soft"]]),
                ("BOX", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("INNERGRID", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return table


def _build_iso45001_document_control_attention_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    controls = list(summary.get("document_controls", []))
    controls.sort(key=lambda control: (_document_control_status(control)[0], control.get("orden", 0), control.get("codigo", "")))
    rows = [["Código", "Control / tipo", "Puntos", "Archivos", "Estado", "Lectura"]]
    attention = [control for control in controls if _document_control_status(control)[0] < 4]
    if not attention:
        rows.append(["-", "Sin controles pendientes.", "-", "-", "Sustentado", "La cobertura documental está completa."])
    for control in attention[:12]:
        _order, status, reading = _document_control_status(control)
        rows.append(
            [
                control["codigo"],
                f"{control['apartado']} {control['nombre']}<br/>{control['clasificacion']}",
                f"{control['puntos_cubiertos']}/{control['total_puntos']}",
                ", ".join(item["archivo_nombre_original"] for item in control["evidencias"]) or "Sin archivo",
                status,
                reading,
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.55 * inch, 2.2 * inch, 0.56 * inch, 1.23 * inch, 1.05 * inch, 1.56 * inch],
        center_from_col=2,
    )


def _build_iso45001_document_control_appendix(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Cód.", "Tipo", "Control / estado", "Puntos", "Archivo(s)", "Reactivos", "Observación", "Contenido mínimo"]]
    for control in summary.get("document_controls", []):
        _order, status, _reading = _document_control_status(control)
        rows.append(
            [
                control["codigo"],
                control["clasificacion"],
                f"{control['nombre']}<br/>{status}",
                f"{control['puntos_cubiertos']}/{control['total_puntos']}",
                ", ".join(item["archivo_nombre_original"] for item in control["evidencias"]) or "Sin archivo",
                ", ".join(item["codigo"] for item in control["reactivos"]) or "-",
                control["observacion"] or "Sin observación",
                control["contenido_minimo"],
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[0.42 * inch, 0.72 * inch, 1.35 * inch, 0.48 * inch, 1.1 * inch, 1.05 * inch, 1.1 * inch, 1.0 * inch],
        center_from_col=3,
    )


def _build_iso45001_document_traceability_appendix(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Archivo", "Control", "Estado", "Puntos", "Reactivos cubiertos", "Observación"]]
    for control in summary.get("document_controls", []):
        files = control["evidencias"] or [None]
        for evidence in files:
            rows.append(
                [
                    evidence["archivo_nombre_original"] if evidence else "Sin archivo vinculado",
                    f"{control['codigo']} - {control['nombre']}",
                    control["estado_label"],
                    f"{control['puntos_cubiertos']}/{control['total_puntos']}",
                    ", ".join(item["codigo"] for item in control["reactivos"]) or "-",
                    control["observacion"] or "Sin observación",
                ]
            )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[1.25 * inch, 1.72 * inch, 0.8 * inch, 0.55 * inch, 1.45 * inch, 1.5 * inch],
        center_from_col=2,
    )


def _build_iso45001_document_coverage_by_clause_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    rows = [["Cláusula", "Controles", "Evaluados", "Puntos", "Archivos", "Brechas", "Lectura"]]
    for clause in summary["clauses"]:
        stats = _document_control_scope_stats(_document_controls_for_clause(summary, clause["numero"]))
        reading = (
            "Sin controles documentales asociados."
            if not stats["total"]
            else (
                "Cobertura sustentada."
                if stats["sustained"] == stats["total"] and stats["total"]
                else f"{stats['gaps']} brecha(s); {stats['sustained']} control(es) con soporte."
            )
        )
        rows.append(
            [
                f"{clause['numero']} {clause['nombre']}",
                str(stats["total"]),
                f"{stats['evaluated']}/{stats['total']}",
                f"{stats['covered_points']}/{stats['total_points']}",
                str(stats["files"]),
                str(stats["gaps"]),
                reading,
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[2.25 * inch, 0.6 * inch, 0.75 * inch, 0.78 * inch, 0.62 * inch, 0.55 * inch, 1.5 * inch],
        center_from_col=1,
    )


def _build_iso45001_clause_heatmap(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    headers = ["Cl\u00e1usula", "0%", "1-20", "21-40", "41-60", "61-80", "81-100", "Cumpl."]
    rows = [[Paragraph(header, styles["table_header"]) for header in headers]]
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), PDF_THEME["primary"]),
        ("GRID", (0, 0), (-1, -1), 0.45, PDF_THEME["line"]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for row_index, clause in enumerate(summary["clauses"], start=1):
        bucket = _iso45001_maturity_bucket(clause["percent"])
        cells = [
            Paragraph(f"{_paragraph_escape(clause['numero'])} {_paragraph_escape(clause['nombre'])}", styles["small"]),
            "", "", "", "", "", "",
            Paragraph(_format_iso45001_percent(clause["percent"]), styles["small_center"]),
        ]
        if bucket is not None:
            color = _iso45001_traffic_color(clause["percent"])
            text_style = styles["traffic_light"] if bucket in {0, 1, 2, 5} else styles["traffic_dark"]
            cells[bucket + 1] = Paragraph("OK", text_style)
            commands.extend(
                [
                    ("BACKGROUND", (bucket + 1, row_index), (bucket + 1, row_index), color),
                    ("TEXTCOLOR", (bucket + 1, row_index), (bucket + 1, row_index), colors.white),
                ]
            )
        rows.append(cells)
    table = Table(
        rows,
        colWidths=[2.45 * inch, 0.52 * inch, 0.58 * inch, 0.62 * inch, 0.62 * inch, 0.62 * inch, 0.7 * inch, 0.72 * inch],
        repeatRows=1,
    )
    table.setStyle(TableStyle(commands))
    return table


def _build_iso45001_priority_sections_table(sections: list[dict], styles: dict[str, ParagraphStyle]) -> Table:
    uses_document_controls = any(
        section.get("document_control_coverage") is not None
        for section in sections
    )
    support_header = "Doc." if uses_document_controls else "Adj."
    rows = [["Apartado", "Avance", "Cumpl.", "No", "Parcial", support_header, "Prioridad"]]
    if not sections:
        rows.append(["Sin apartados evaluables para priorizar", "-", "-", "-", "-", "-", "-"])
    for section in sections[:10]:
        counts = _iso45001_section_response_counts(section)
        coverage = section["evidence_coverage"]
        document_coverage = section.get("document_control_coverage")
        support_value = (
            (
                f"{document_coverage['sustained_controls']}/"
                f"{document_coverage['total_controls']}"
                if document_coverage["total_controls"]
                else "-"
            )
            if document_coverage is not None
            else f"{coverage['satisfied']}/{coverage['required']}"
        )
        rows.append(
            [
                f"{section['codigo']} {section['nombre']}",
                _format_iso45001_percent(section["completion"]),
                _format_iso45001_percent(section["percent"]),
                str(counts["no"]),
                str(counts["parcial"]),
                support_value,
                _iso45001_priority_label(section["percent"], section["completion"]),
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[2.6 * inch, 0.7 * inch, 0.72 * inch, 0.48 * inch, 0.62 * inch, 0.7 * inch, 1.0 * inch],
        center_from_col=1,
    )


def _build_iso45001_evidence_coverage_table(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        return _build_iso45001_document_coverage_by_clause_table(summary, styles)
    rows = [["Cl\u00e1usula", "Archivos", "Adjuntos exigidos", "Pend.", "Lectura"]]
    for clause in summary["clauses"]:
        coverage = clause["evidence_coverage"]
        rows.append(
            [
                f"{clause['numero']} {clause['nombre']}",
                str(clause["evidence_count"]),
                f"{coverage['satisfied']}/{coverage['required']}",
                str(coverage["missing"]),
                _iso45001_evidence_reading(coverage),
            ]
        )
    return _iso45001_pdf_table(
        rows,
        styles,
        col_widths=[2.55 * inch, 0.72 * inch, 1.05 * inch, 0.58 * inch, 2.35 * inch],
        center_from_col=1,
    )


def _build_iso45001_radar_chart(summary: dict) -> Drawing:
    clauses = summary["clauses"]
    drawing = Drawing(520, 292)
    center_x, center_y, radius = 225, 133, 92
    label_radius = radius + 27
    drawing.add(String(260, 280, "Ara\u00f1a de madurez por cl\u00e1usula", fontName=ISO45001_PDF_FONT_BOLD, fontSize=11, fillColor=PDF_THEME["primary"], textAnchor="middle"))
    drawing.add(String(260, 264, "\u00cdndice comparativo de cumplimiento 0-100; todos los reactivos son aplicables", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=7.8, fillColor=PDF_THEME["muted"], textAnchor="middle"))
    if not clauses:
        drawing.add(String(center_x, center_y, "Sin cl\u00e1usulas disponibles", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=8.5, fillColor=PDF_THEME["muted"], textAnchor="middle"))
        return drawing
    angles = [math.pi / 2 - (2 * math.pi * index / len(clauses)) for index in range(len(clauses))]
    for level in (25, 50, 75, 100):
        points = []
        for angle in angles:
            points.extend([center_x + math.cos(angle) * radius * level / 100, center_y + math.sin(angle) * radius * level / 100])
        drawing.add(Polygon(points, fillColor=None, strokeColor=PDF_THEME["line"], strokeWidth=0.65))
        drawing.add(String(center_x + 13, center_y + radius * level / 100 - 2, str(level), fontName=ISO45001_PDF_FONT_BOLD, fontSize=6.5, fillColor=PDF_THEME["muted"]))
    value_points = []
    for angle, clause in zip(angles, clauses):
        axis_x = center_x + math.cos(angle) * radius
        axis_y = center_y + math.sin(angle) * radius
        drawing.add(Line(center_x, center_y, axis_x, axis_y, strokeColor=PDF_THEME["line"], strokeWidth=0.5))
        label_x = center_x + math.cos(angle) * label_radius
        label_y = center_y + math.sin(angle) * label_radius
        # The lower two labels sit close together in a seven-axis radar.  Give
        # them a small outward offset so C7 and C8 remain independently
        # legible in a printed executive report.
        if clause["numero"] == "7":
            label_x += 10
            label_y -= 5
        elif clause["numero"] == "8":
            label_x -= 10
            label_y -= 5
        anchor = "middle"
        if label_x < center_x - 8:
            anchor = "end"
        elif label_x > center_x + 8:
            anchor = "start"
        drawing.add(String(label_x, label_y - 3, f"C{clause['numero']}", fontName=ISO45001_PDF_FONT_BOLD, fontSize=7.7, fillColor=PDF_THEME["ink"], textAnchor=anchor))
        percent = min(clause["percent"] or 0, 100)
        value_points.extend([center_x + math.cos(angle) * radius * percent / 100, center_y + math.sin(angle) * radius * percent / 100])
    if any((clause["percent"] or 0) > 0 for clause in clauses):
        color = _iso45001_traffic_color(summary["percent"])
        drawing.add(Polygon(value_points, fillColor=_iso45001_traffic_soft_color(summary["percent"]), strokeColor=color, strokeWidth=1.8))
        for index in range(0, len(value_points), 2):
            drawing.add(Rect(value_points[index] - 2, value_points[index + 1] - 2, 4, 4, fillColor=color, strokeColor=color))
    else:
        drawing.add(String(center_x, center_y - 4, "Sin cumplimiento registrado", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=8.2, fillColor=PDF_THEME["muted"], textAnchor="middle"))
    legend = [
        ("0-40", "Atenci\u00f3n prioritaria", PDF_THEME["red"]),
        ("41-80", "Implementaci\u00f3n parcial", PDF_THEME["yellow"]),
        ("81-100", "Madurez alta", PDF_THEME["green"]),
    ]
    for index, (range_label, label, color) in enumerate(legend):
        y = 152 - index * 18
        drawing.add(Rect(387, y, 8, 8, fillColor=color, strokeColor=color))
        drawing.add(String(401, y + 1, f"{range_label}: {label}", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=7.2, fillColor=PDF_THEME["muted"]))
    return drawing


def _build_iso45001_metric_cards(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        return _build_iso45001_v2_metric_cards(summary, styles)
    counts = summary["response_counts"]
    evidence = summary["evidence_coverage"]
    cards = [
        ("Avance", _format_iso45001_percent(summary["completion"]), f"{summary['answered_questions']} de {summary['total_questions']} reactivos"),
        ("Cumplimiento", _format_iso45001_percent(summary["percent"]), "Puntos sobre el m\u00e1ximo posible"),
        ("Madurez", _compact_iso45001_maturity(summary["maturity_label"]), _maturity_level_prefix_iso45001(summary["maturity_label"])),
        ("Evidencias", str(summary["evidence_count"]), "Archivos activos"),
        ("Brechas No", str(counts["no"]), "Atenci\u00f3n prioritaria"),
        ("Brechas Parcial", str(counts["parcial"]), "Implementaci\u00f3n por completar"),
        ("Sin respuesta", str(counts["sin_respuesta"]), "Pendientes de captura"),
        ("Adjuntos exigidos", f"{evidence['satisfied']}/{evidence['required']}", "Soporte documental requerido"),
    ]
    rows = []
    for start in range(0, len(cards), 4):
        row = []
        for label, value, caption in cards[start : start + 4]:
            card = Table(
                [
                    [Paragraph(_paragraph_escape(label.upper()), styles["metric_label"])],
                    [Paragraph(_paragraph_escape(value), styles["metric_value"])],
                    [Paragraph(_paragraph_escape(caption), styles["metric_caption"])],
                ],
                colWidths=[1.55 * inch],
                rowHeights=[0.22 * inch, 0.42 * inch, 0.28 * inch],
            )
            card.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 2),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
            row.append(card)
        rows.append(row)
    table = Table(rows, colWidths=[1.76 * inch] * 4, rowHeights=[0.98 * inch] * len(rows))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, PDF_THEME["primary_soft"]]),
                ("BOX", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("INNERGRID", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return table


def _build_iso45001_methodology_panel(summary: dict, styles: dict[str, ParagraphStyle]) -> Table:
    if _uses_document_control_coverage(summary):
        stats = summary["document_evidence_stats"]
        method = Paragraph(
            "Método de cálculo ISO: No=0, Parcial=1 y Sí=2; todos los reactivos son aplicables. "
            "Cumplimiento = puntos obtenidos / (reactivos x 2) y avance = respuestas capturadas. "
            f"La cobertura documental se muestra aparte: {stats['covered_points']} de {stats['total_points']} puntos, "
            f"{stats['evaluated_controls']} de {stats['total_controls']} controles evaluados y "
            f"{stats['file_count']} archivo(s) reutilizable(s).",
            styles["body"],
        )
        layout = Table([[method, _build_iso45001_maturity_legend(styles)]], colWidths=[4.18 * inch, 2.85 * inch])
        layout.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                    ("BOX", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                    ("INNERGRID", (0, 0), (-1, -1), 0.45, PDF_THEME["line"]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        return layout
    evidence = summary["evidence_coverage"]
    method = Paragraph(
        "M\u00e9todo de c\u00e1lculo: No=0, Parcial=1 y S\u00ed=2. Todos los reactivos del cat\u00e1logo son aplicables; "
        "Cumplimiento = puntos obtenidos / (reactivos x 2). Avance refleja respuestas capturadas. "
        f"Adjuntos exigidos satisfechos: {evidence['satisfied']} de {evidence['required']}.",
        styles["body"],
    )
    layout = Table([[method, _build_iso45001_maturity_legend(styles)]], colWidths=[4.18 * inch, 2.85 * inch])
    layout.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.55, PDF_THEME["line"]),
                ("INNERGRID", (0, 0), (-1, -1), 0.45, PDF_THEME["line"]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return layout


def _build_iso45001_maturity_legend(styles: dict[str, ParagraphStyle]) -> Table:
    rows = [[Paragraph("Rango", styles["table_header"]), Paragraph("Nivel", styles["table_header"])]]
    for range_label, label, color_slug in MATURITY_REPORT_GUIDE:
        label_style = styles["traffic_light"] if color_slug in {"red", "green"} else styles["traffic_dark"]
        rows.append([Paragraph(range_label, label_style), Paragraph(label, styles["small"])])
    table = Table(rows, colWidths=[0.65 * inch, 1.85 * inch], repeatRows=1)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), PDF_THEME["primary"]),
        ("GRID", (0, 0), (-1, -1), 0.4, PDF_THEME["line"]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for index, (_range, _label, color_slug) in enumerate(MATURITY_REPORT_GUIDE, start=1):
        commands.append(("BACKGROUND", (0, index), (0, index), PDF_THEME[color_slug]))
    table.setStyle(TableStyle(commands))
    return table


def _build_iso45001_chart_pair(summary: dict) -> Table:
    table = Table(
        [[_build_iso45001_clause_bar_chart(summary), _build_iso45001_response_distribution_chart(summary)]],
        colWidths=[4.65 * inch, 2.6 * inch],
    )
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    return table


def _build_iso45001_clause_bar_chart(summary: dict) -> Drawing:
    drawing = Drawing(330, 205)
    drawing.add(String(165, 188, "Cumplimiento por cl\u00e1usula", fontName=ISO45001_PDF_FONT_BOLD, fontSize=10, fillColor=PDF_THEME["primary"], textAnchor="middle"))
    plot_x, plot_y, plot_width, plot_height = 36, 36, 268, 122
    drawing.add(Rect(plot_x, plot_y, plot_width, plot_height, fillColor=colors.white, strokeColor=PDF_THEME["line"], strokeWidth=0.8))
    for tick in range(5):
        value = 25 * tick
        y = plot_y + (plot_height * value / 100)
        drawing.add(Line(plot_x, y, plot_x + plot_width, y, strokeColor=PDF_THEME["line"], strokeWidth=0.45))
        drawing.add(String(plot_x - 7, y - 2, str(value), fontName=ISO45001_PDF_FONT_REGULAR, fontSize=6.5, fillColor=PDF_THEME["muted"], textAnchor="end"))
    clauses = summary["clauses"]
    slot = plot_width / max(len(clauses), 1)
    bar_width = min(20, slot * 0.55)
    for index, clause in enumerate(clauses):
        percent = clause["percent"] or 0
        x = plot_x + slot * index + (slot - bar_width) / 2
        height = plot_height * min(percent, 100) / 100
        color = _iso45001_traffic_color(clause["percent"])
        drawing.add(Rect(x, plot_y, bar_width, height, fillColor=color, strokeColor=color, strokeWidth=0))
        drawing.add(String(x + bar_width / 2, plot_y + height + 5, _format_iso45001_percent(clause["percent"]), fontName=ISO45001_PDF_FONT_BOLD, fontSize=6.2, fillColor=PDF_THEME["primary_dark"], textAnchor="middle"))
        drawing.add(String(x + bar_width / 2, plot_y - 13, f"C{clause['numero']}", fontName=ISO45001_PDF_FONT_BOLD, fontSize=7.1, fillColor=PDF_THEME["muted"], textAnchor="middle"))
    drawing.add(String(plot_x + plot_width / 2, 10, "Cl\u00e1usulas auditables 4 a 10", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=7.2, fillColor=PDF_THEME["muted"], textAnchor="middle"))
    return drawing


def _build_iso45001_response_distribution_chart(summary: dict) -> Drawing:
    rows = _iso45001_response_distribution(summary)
    total = sum(row["count"] for row in rows)
    drawing = Drawing(185, 205)
    drawing.add(String(92, 188, "Distribuci\u00f3n de respuestas", fontName=ISO45001_PDF_FONT_BOLD, fontSize=10, fillColor=PDF_THEME["primary"], textAnchor="middle"))
    bar_x, bar_y, bar_width, bar_height = 14, 150, 156, 16
    drawing.add(Rect(bar_x, bar_y, bar_width, bar_height, fillColor=PDF_THEME["primary_soft"], strokeColor=PDF_THEME["line"], strokeWidth=0.7))
    if total <= 0:
        drawing.add(String(92, 112, "Sin respuestas registradas", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=8.5, fillColor=PDF_THEME["muted"], textAnchor="middle"))
        return drawing
    offset = bar_x
    for row in rows:
        width = bar_width * row["count"] / total
        if width > 0:
            drawing.add(Rect(offset, bar_y, width, bar_height, fillColor=row["color"], strokeColor=row["color"], strokeWidth=0))
        offset += width
    for index, row in enumerate(rows):
        y = 121 - index * 23
        drawing.add(Rect(14, y, 9, 9, fillColor=row["color"], strokeColor=row["color"]))
        drawing.add(String(30, y + 1, row["label"], fontName=ISO45001_PDF_FONT_BOLD, fontSize=7.4, fillColor=PDF_THEME["primary_dark"]))
        drawing.add(String(171, y + 1, f"{row['count']} | {_format_iso45001_percent(row['percent'])}", fontName=ISO45001_PDF_FONT_REGULAR, fontSize=7.2, fillColor=PDF_THEME["muted"], textAnchor="end"))
    return drawing


def _build_iso45001_progress_chart(summary: dict) -> Drawing:
    drawing = Drawing(520, 72)
    drawing.add(String(2, 58, "Avance de captura vs. cumplimiento", fontName=ISO45001_PDF_FONT_BOLD, fontSize=9.5, fillColor=PDF_THEME["primary"], textAnchor="start"))
    rows = [
        ("Avance de captura", summary["completion"], _iso45001_traffic_color(summary["completion"])),
        ("Cumplimiento ISO", summary["percent"] or 0, _iso45001_traffic_color(summary["percent"])),
    ]
    for index, (label, value, color) in enumerate(rows):
        y = 35 - index * 23
        drawing.add(String(4, y + 2, label, fontName=ISO45001_PDF_FONT_REGULAR, fontSize=7.8, fillColor=PDF_THEME["muted"]))
        drawing.add(Rect(120, y, 330, 11, fillColor=PDF_THEME["primary_soft"], strokeColor=PDF_THEME["line"], strokeWidth=0.4))
        if value > 0:
            drawing.add(Rect(120, y, 330 * min(value, 100) / 100, 11, fillColor=color, strokeColor=color, strokeWidth=0))
        drawing.add(String(463, y + 2, _format_iso45001_percent(value), fontName=ISO45001_PDF_FONT_BOLD, fontSize=7.6, fillColor=PDF_THEME["primary_dark"]))
    return drawing
